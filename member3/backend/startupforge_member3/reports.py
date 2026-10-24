"""Formal, source-linked report content and DOCX/PDF exports."""

from __future__ import annotations

import io
import os
import re
from pathlib import Path
from typing import Annotated, Any

from pydantic import BaseModel, ConfigDict, Field, ValidationError

from app.integration.contracts import GenerationResult, TemplateDescriptor
from app.integration.interfaces import ReportsAdapter as ReportsProtocol
from app.modules.intelligence.contracts import ProjectContext
from app.modules.intelligence.errors import IntelligenceError
from app.modules.intelligence.output_author import OutputAuthor

from .core.artifacts import ArtifactStore


class ContentModel(BaseModel):
    model_config = ConfigDict(extra="forbid")


class ReportSection(ContentModel):
    section_id: str = Field(min_length=1, max_length=100)
    heading: str = Field(min_length=1, max_length=200)
    paragraphs: list[str] = Field(min_length=1, max_length=40)
    source_ids: list[str] = Field(default_factory=list)
    evidence_ids: list[str] = Field(default_factory=list)
    claim_kinds: list[str] = Field(default_factory=list)


class ReportReference(ContentModel):
    source_id: str
    title: str
    url: str
    publisher: str | None = None
    retrieved_at: str


class ReportContent(ContentModel):
    title: str = Field(min_length=1, max_length=200)
    sections: list[ReportSection] = Field(min_length=1, max_length=40)
    bibliography: list[ReportReference] = Field(default_factory=list, max_length=100)


class ReportNarrativeSection(ContentModel):
    section_id: str = Field(min_length=1, max_length=100)
    paragraphs: list[Annotated[str, Field(min_length=100, max_length=1800)]] = Field(min_length=2, max_length=3)


class ReportNarrative(ContentModel):
    sections: list[ReportNarrativeSection] = Field(min_length=1, max_length=3)


_REPORT_TITLE_SUFFIX = " — Startup Feasibility Report"
_REPORT_TITLE_MAX_LENGTH = 200


def _recommendations_for_gaps(gaps: list[str]) -> list[str]:
    """Turn unresolved questions into proposed tests without presenting them as results."""
    tracks: dict[str, str] = {}
    for gap in gaps:
        lowered = gap.lower()
        if any(term in lowered for term in ("willingness to pay", "pricing", "price", "cost", "financial", "revenue", "margin")):
            tracks["economics"] = ("Economics: gather current customer price expectations and supplier or delivery costs, then show the assumptions in a simple per-customer model. Do not report a forecast until its inputs have been checked.")
        elif any(term in lowered for term in ("demand", "customer", "student", "user", "adoption", "need", "interest")):
            tracks["customer"] = ("Customer need: speak with people who fit the stated target group about their current behavior, then test the proposed offer with a concrete action that goes beyond stated interest.")
        elif any(term in lowered for term in ("market", "competitor", "alternative", "substitute", "segment")):
            tracks["market"] = ("Market and alternatives: define the first reachable customer group and region, identify the options those customers use now, and source each input used in a market estimate.")
        elif any(term in lowered for term in ("operational", "operation", "scalability", "scale", "capacity", "pickup", "fulfillment", "partner", "delivery")):
            tracks["operations"] = ("Operations: run a limited end-to-end test and record completion, delays, support needs, capacity, and partner dependencies before expanding the service.")
        elif any(term in lowered for term in ("source", "evidence", "citation", "fetch", "download", "off-topic", "off topic")):
            tracks["evidence"] = ("Evidence quality: obtain a relevant, accessible source for each material claim, retain the exact passage and publication details, and remove inaccessible or off-topic pages from the evidence base.")
        else:
            tracks.setdefault("general", "Remaining question: state what evidence would resolve it, choose a practical way to collect that evidence, and record the result before making a decision.")
    if not tracks:
        tracks = {
            "customer": "Customer need: confirm the target users' current behavior and test the offer with an observable action.",
            "operations": "Operations: check whether the core offer can be delivered reliably and identify any partner or capacity limits.",
            "economics": "Economics: document who pays, what it costs to deliver, and which inputs are observed versus estimated.",
        }
    return ["The actions below are proposed ways to resolve the recorded gaps. They are recommendations, not completed research or pilot results.", *tracks.values()]


def _qualify_unknown_claim(heading: str, text: str, kind: str) -> str:
    """Replace repeated reviewer boilerplate with a section-specific unknown."""
    if kind != "unknown" or not text.startswith("This conclusion remains unknown"):
        return text
    lowered_heading = heading.lower()
    lowered_text = text.lower()
    if "alternative" in lowered_heading and "critic's concern" in lowered_text:
        return "The review concern remains open; compare the proposed offer with alternatives using evidence from intended customers."
    if "alternative" in lowered_heading:
        return "Available sources do not establish how the proposed offer compares with alternatives used by intended customers."
    if "feasibility" in lowered_heading:
        return "The context does not establish whether the proposed solution can be delivered within the stated constraints."
    if "risk" in lowered_heading:
        return "Available evidence does not resolve which risks could prevent the proposed solution from serving intended users."
    if "assumption" in lowered_heading:
        return "The founder's assumptions remain untested with direct evidence from intended users."
    return text


def _report_title(project_name: str) -> str:
    name = re.sub(r"\s+", " ", project_name).strip() or "Startup"
    name_limit = _REPORT_TITLE_MAX_LENGTH - len(_REPORT_TITLE_SUFFIX)
    if len(name) > name_limit:
        shortened = name[:name_limit - 1]
        boundary = shortened.rfind(" ")
        if boundary >= name_limit * 0.7:
            shortened = shortened[:boundary]
        name = shortened.rstrip(" ,;:—-") + "…"
    return f"{name}{_REPORT_TITLE_SUFFIX}"


class ReportsAdapter(ReportsProtocol):
    def __init__(self, artifacts: ArtifactStore, author: OutputAuthor | None = None):
        self.artifacts = artifacts
        self.author = author

    def templates(self) -> list[TemplateDescriptor]:
        return [TemplateDescriptor(template_id="evidence-led-report", name="Evidence-led startup report", kind="report",
                                   supported_formats=["docx", "pdf"],
                                   supported_options={"sections": ["Executive summary", "Idea and target users",
                                                                      "Product and service concept", "Market scope and alternatives",
                                                                      "Business model and financial questions", "Pilot plan and measures",
                                                                      "Research scope and limitations", "Open questions and evidence gaps",
                                                                      "Assessment and decision", "Sources"]})]

    def validate_options(self, kind: str, options: dict[str, Any]) -> None:
        if kind != "report" or options.get("template_id") != "evidence-led-report":
            raise IntelligenceError("TEMPLATE_NOT_FOUND", "Choose an installed report template.", 422)
        formats = options.get("formats", [])
        if not formats or any(fmt not in {"docx", "pdf"} for fmt in formats):
            raise IntelligenceError("UNSUPPORTED_FORMAT", "Report exports support DOCX and PDF.", 422,
                                    details={"supported_formats": ["docx", "pdf"]})

    def validate_content(self, content: dict[str, Any], context: ProjectContext) -> dict[str, Any]:
        try:
            normalized = ReportContent.model_validate(content).model_dump(mode="json")
        except ValidationError as exc:
            raise IntelligenceError("INVALID_ARTIFACT_CONTENT", "The edited content does not match the report format.", 422,
                                    details={"issues": exc.errors(include_input=False)}) from exc
        source_ids = {source.source_id for source in context.sources}
        evidence_by_id = {item.evidence_id: item for item in context.evidence}
        bibliography_ids = {item["source_id"] for item in normalized["bibliography"]}
        if bibliography_ids - source_ids:
            raise IntelligenceError("INVALID_CONTENT_REFERENCE", "The bibliography contains a source outside this context.", 422)
        for section in normalized["sections"]:
            if set(section["source_ids"]) - source_ids or set(section["evidence_ids"]) - set(evidence_by_id):
                raise IntelligenceError("INVALID_CONTENT_REFERENCE", "Report content references material outside this context.", 422)
            if any(evidence_by_id[evidence_id].source_id not in section["source_ids"] for evidence_id in section["evidence_ids"]):
                raise IntelligenceError("INVALID_CONTENT_REFERENCE", "Each cited evidence item needs its matching source reference.", 422)
            if "supported" in section["claim_kinds"] and not section["evidence_ids"]:
                raise IntelligenceError("UNSOURCED_SUPPORTED_CLAIM", "Supported report findings must keep their evidence references.", 422)
            for paragraph in section["paragraphs"]:
                if re.search(r"(?<![A-Za-z])(?:\$|€|£|₹|INR\s*)?\d[\d,.]*(?:\s*%|\s*(?:million|billion|thousand|crore|lakh))?", paragraph, re.I):
                    if not section["evidence_ids"] and not ({"assumption", "provided", "reviewed"} & set(section["claim_kinds"])):
                        raise IntelligenceError("UNSOURCED_NUMERIC_CLAIM", "Numeric report claims need evidence or an explicit assumption label.", 422)
        return normalized

    def build_content(self, context: ProjectContext) -> dict[str, Any]:
        idea = context.idea
        sections: list[ReportSection] = []

        def add(section_id: str, heading: str, paragraphs: list[str], source_ids: list[str] | None = None,
                evidence_ids: list[str] | None = None, claim_kinds: list[str] | None = None) -> None:
            if paragraphs:
                sections.append(ReportSection(section_id=section_id, heading=heading, paragraphs=paragraphs,
                                              source_ids=sorted(set(source_ids or [])),
                                              evidence_ids=sorted(set(evidence_ids or [])),
                                              claim_kinds=claim_kinds or []))

        add("executive-summary", "Executive summary", [context.final_analysis.summary], claim_kinds=["reviewed"])
        add("idea-and-users", "Idea and target users", [f"Problem: {idea.problem}", f"Proposed solution: {idea.solution}",
             "Target users: " + (", ".join(idea.target_users) if idea.target_users else "Not specified in the submitted idea."),
             f"Domain: {idea.domain}"], claim_kinds=["provided"])
        if idea.assumptions:
            add("founder-assumptions", "Founder assumptions", [f"Assumption: {item}" for item in idea.assumptions], claim_kinds=["assumption"])

        for analysis_section in context.final_analysis.sections:
            paragraphs: list[str] = []
            source_ids: set[str] = set()
            evidence_ids: set[str] = set()
            claim_kinds: set[str] = set()
            for claim in analysis_section.claims:
                label = {"supported": "Supported finding", "assumption": "Unverified assumption", "unknown": "Unknown"}[claim.kind]
                claim_text = _qualify_unknown_claim(analysis_section.heading, claim.text, claim.kind)
                paragraphs.append(f"{label}: {claim_text}")
                evidence_ids.update(claim.evidence_ids)
                claim_kinds.add(claim.kind)
            source_ids.update(e.source_id for e in context.evidence if e.evidence_id in evidence_ids)
            add(f"analysis-{analysis_section.section_id}", analysis_section.heading, paragraphs,
                list(source_ids), list(evidence_ids), sorted(claim_kinds))

        if context.figures:
            paragraphs = [f"{figure.display_value} — {figure.metric}\nQuoted evidence: {figure.excerpt}\nQualification: {figure.qualification}"
                          for figure in context.figures]
            add("sourced-statistics", "Sourced statistics and market figures", paragraphs,
                [figure.source_id for figure in context.figures], [figure.evidence_id for figure in context.figures], ["supported"])

        source_count = len(context.sources)
        evidence_count = len(context.evidence)
        figure_count = len(context.figures)
        add("product-concept", "Product and service concept", [
            f"The proposed offer is: {idea.solution}",
            "A first prototype should test the smallest complete user task that demonstrates this offer. Record where users hesitate, fail, or need help, and revise the workflow before adding features. Compare this proposed approach with any prototype results already recorded in the project context.",
        ], claim_kinds=["provided", "recommendation"])

        audience = ", ".join(idea.target_users) if idea.target_users else "the user groups identified during discovery"
        add("customer-validation", "Customer need and validation", [
            f"The first customer conversations should focus on {audience}. Ask people to describe how they handle the problem today, how often it occurs, and what they have already tried. Compare those behaviors with the stated problem: {idea.problem}",
            "Treat stated interest as an early signal only. Show a concrete version of the offer and look for an observable next step, such as trying the workflow, requesting access, or making a purchase when payment is part of the idea. Review any existing customer test results and collect further evidence for the questions they leave unresolved.",
        ], claim_kinds=["provided", "recommendation"])

        market_paragraphs = [
            f"The saved context contains {figure_count} extracted figures, {source_count} sources, and {evidence_count} evidence excerpts. These counts describe the research record; they do not establish the size of the opportunity or demand for this specific product."
        ]
        if figure_count == 0:
            market_paragraphs.append("The saved context has no separately extracted statistics. Check any numerical claims in the analysis for their source and scope, and obtain directly relevant evidence before using them to estimate market size, growth, or willingness to pay.")
        market_paragraphs.append("A useful market estimate should define the first reachable customer group, region, and use or purchase frequency. Build the estimate from traceable inputs and compare the offer with the alternatives those customers actually use. Broad industry totals can provide background, but they do not show how many customers this project can reach.")
        add("market-scope", "Market scope and alternatives", market_paragraphs, claim_kinds=["reviewed", "recommendation"])

        assumption_text = "; ".join(idea.assumptions) if idea.assumptions else "No founder assumptions were entered in the saved idea."
        add("business-model", "Business model and financial questions", [
            f"The assumptions recorded by the founder are: {assumption_text}",
            "Before preparing a financial forecast, specify who pays, what the customer pays for, how often the payment occurs, and what it costs to deliver the offer. Validate each major input with customer behavior, supplier or partner quotes, and a clearly stated calculation. Review any recorded prices, costs, revenue, or margins for their source, date, and relevance to this project, and gather evidence for missing inputs.",
            "Keep the business model as a hypothesis until customers demonstrate a credible commitment and delivery costs are understood. Separate measured inputs from estimates in any later forecast.",
        ], claim_kinds=["provided", "recommendation"])

        add("pilot-plan", "Pilot plan and measures", [
            "Begin with a limited test that uses the intended customer group and one complete version of the core offer. Document how participants are recruited, what they are asked to do, and which parts of the experience are included.",
            "Track observable measures that fit the idea: whether participants complete the core task, where they drop out, what support they need, whether they return, and whether they take the next step the business model requires. Set the success criteria before the test begins, distinguish proposed thresholds from measured outcomes, and use any existing pilot evidence to identify the next unresolved question.",
            "For any external partner or operational dependency, confirm availability, responsibilities, capacity, and service expectations before inviting customers into the pilot. Use what the test shows to decide whether to revise the offer, gather more evidence, or proceed to a broader trial.",
        ], claim_kinds=["recommendation"])

        constraints = [item.strip() for item in idea.constraints if item.strip()]
        critic_notes = [f"{issue.severity.title()} review issue: {issue.reason}"
                        + (f"\nResolution recorded: {issue.resolution}" if issue.resolution else "")
                        for issue in context.critic.issues]
        risk_paragraphs = []
        if constraints:
            risk_paragraphs.append("Founder-stated constraints:\n" + "\n".join(constraints))
        if critic_notes:
            risk_paragraphs.extend(critic_notes)
        if not risk_paragraphs:
            risk_paragraphs.append("The saved context contains no explicit founder constraints or critic issues. This means they were not recorded in this review; it does not establish that operational, regulatory, privacy, or partner risks have been checked.")
        risk_paragraphs.append("Before a pilot, identify the dependencies that could prevent the core offer from reaching the intended user. Assign an owner and a practical check for each dependency, and record any unresolved issue as a launch condition rather than treating it as settled.")
        if len(risk_paragraphs) > 40:
            # Consolidate adjacent notes without dropping constraints, issues, or resolutions.
            group_size = (len(risk_paragraphs) + 39) // 40
            risk_paragraphs = ["\n\n".join(risk_paragraphs[index:index + group_size])
                               for index in range(0, len(risk_paragraphs), group_size)]
        add("risks-and-constraints", "Risks, constraints, and dependencies", risk_paragraphs,
            claim_kinds=["provided", "reviewed", "recommendation"])

        warning_text = " ".join(context.warnings).strip()
        revision_label = "revision pass" if context.execution.revision_passes == 1 else "revision passes"
        research_method = (f"The saved research run used {context.execution.mode} research, {context.execution.search_provider} search, "
                           f"and {context.execution.retrieval_method} retrieval, with {context.execution.revision_passes} {revision_label}. "
                           f"It produced {source_count} sources, {evidence_count} evidence excerpts, and {figure_count} extracted figures.")
        limitations = [research_method,
                       "The final analysis and critic review help identify claims that are supported, assumed, or still unknown. A source excerpt supports only the point it actually states; it does not by itself validate demand, market size, feasibility, or financial performance for this project."]
        if warning_text:
            lowered_warnings = warning_text.lower()
            if "demo mode" in lowered_warnings or "synthetic fixture" in lowered_warnings:
                limitations.append("The saved context is a demonstration using synthetic fixture sources, not live research or a validated assessment of this startup.")
            if "bounded set" in lowered_warnings or "at most six" in lowered_warnings:
                limitations.append("Web research is limited to a bounded set of fetched public pages and selected passages. The original publications should be checked for their dates, geography, methods, and full context.")
            if "off-topic" in lowered_warnings or "off topic" in lowered_warnings:
                limitations.append("Pages that did not match the idea's core topic were omitted. The topic filter is a conservative lexical check and does not establish semantic relevance.")
            if any(term in lowered_warnings for term in ("page_timeout", "page_access_denied", "no_readable_content", "could not be fetched")):
                limitations.append("Some candidate pages timed out, blocked automated access, or contained too little readable text to review. They were not used as evidence.")
            if "complete literal source quotations" in lowered_warnings or "quotation traceability" in lowered_warnings:
                limitations.append("The evidence guard checks whether supported claims match source wording. That traceability check does not prove a source is accurate or commercially relevant.")
        add("research-scope", "Research scope and limitations", limitations, claim_kinds=["reviewed"])

        gaps = list(dict.fromkeys([*context.final_analysis.gaps, *context.critic.unresolved_gaps]))
        add("open-gaps", "Open questions and evidence gaps", gaps or ["No unresolved gaps were recorded by the review."],
            claim_kinds=["unknown"])
        add("next-steps", "Recommended next steps", _recommendations_for_gaps(gaps), claim_kinds=["recommendation"])
        status = "ready" if context.critic.review_status == "ready" else "ready_with_gaps"
        decision_text = {
            "ready": "The saved review is ready to advance to a scoped pilot. Confirm that the pilot tests the core user need and delivery assumptions, and keep any conclusions within the evidence collected.",
            "ready_with_gaps": "The saved review is ready with gaps. The next decision is whether to run a limited validation test against the unresolved questions. Treat launch, scale, and financial outcomes as unproven until the relevant evidence is collected.",
            "blocked": "The saved review is blocked. Resolve the review issues and obtain the missing evidence before relying on this analysis to make a launch or investment decision.",
        }[status]
        add("decision", "Assessment and decision", [decision_text,
             "Before starting a pilot, write down what result would support proceeding, what result would require a change, and what result would stop the test. The current report makes no claim that these criteria have already been met."],
            claim_kinds=["reviewed", "recommendation"])
        add("review", "Review status", [f"Review status: {status.replace('_', ' ')}.",
             context.critic.revision_summary or "No revision summary was provided."], claim_kinds=["reviewed"])
        bibliography = [ReportReference(source_id=source.source_id, title=source.title, url=source.url,
                                        publisher=source.publisher, retrieved_at=source.retrieved_at.isoformat()).model_dump(mode="json")
                        for source in context.sources]
        return self.validate_content({"title": _report_title(idea.project_name),
                                      "sections": [section.model_dump(mode="json") for section in sections],
                                      "bibliography": bibliography}, context)

    async def _write_narrative(self, content: dict[str, Any], context: ProjectContext, progress) -> dict[str, Any]:
        """Author startup-specific strategy while retaining the immutable research record."""
        requested_ids = {
            "executive-summary", "product-concept", "customer-validation", "market-scope",
            "business-model", "pilot-plan", "next-steps", "decision",
        }
        selected = [section for section in content["sections"] if section["section_id"] in requested_ids]
        idea = context.idea.model_dump(mode="json", exclude={"original_text"})
        gaps = list(dict.fromkeys([*context.final_analysis.gaps, *context.critic.unresolved_gaps]))
        written: dict[str, list[str]] = {}
        for offset in range(0, len(selected), 3):
            batch = selected[offset:offset + 3]
            batch_number = offset // 3 + 1
            await progress(f"ai_writing_report_{batch_number}", 24 + round(offset / len(selected) * 44))
            payload: dict[str, Any] = {
                "idea": idea,
                "reviewed_summary": context.final_analysis.summary,
                "recorded_gaps": gaps,
                "requested_sections": [{"section_id": section["section_id"], "heading": section["heading"]}
                                       for section in batch],
                "writing_guidance": (
                    "Write two substantial paragraphs per section, about seventy to one hundred words each. "
                    "Use startup-specific reasoning and actionable proposals. Keep generated prose free of digits, "
                    "numerical measurements and invented factual claims. Spell out alphanumeric product labels "
                    "in words when needed. The application preserves the exact "
                    "provided idea, reviewed findings, figures, research snapshot and source bibliography separately."
                ),
            }
            expected = {section["section_id"] for section in batch}
            schema = ReportNarrative.model_json_schema()
            schema["properties"]["sections"]["minItems"] = len(batch)
            schema["properties"]["sections"]["maxItems"] = len(batch)
            schema["$defs"]["ReportNarrativeSection"]["properties"]["section_id"]["enum"] = sorted(expected)
            for attempt in range(2):
                try:
                    response = await self.author.generate_json(
                        "report_narrative", payload, schema
                    )
                    narrative = ReportNarrative.model_validate(response)
                    actual = [section.section_id for section in narrative.sections]
                    if len(actual) != len(expected) or set(actual) != expected:
                        raise ValueError("Return each requested section_id exactly once, without extra or missing sections.")
                    for section in narrative.sections:
                        for paragraph in section.paragraphs:
                            if re.search(r"\d", paragraph):
                                raise ValueError("Generated narrative must contain no numerical digits or measurements; exact numbers are retained separately.")
                            if re.search(r"https?://|\b(?:src|ev|evidence|fig)_[A-Za-z0-9_-]+", paragraph, re.I):
                                raise ValueError("Do not include URLs or evidence/source IDs in narrative paragraphs.")
                    written.update({section.section_id: section.paragraphs for section in narrative.sections})
                    break
                except (ValidationError, ValueError, IntelligenceError) as exc:
                    if isinstance(exc, IntelligenceError) and exc.code not in {"MODEL_INVALID_JSON", "PROVIDER_INVALID_RESPONSE"}:
                        raise
                    feedback = ("Return exactly the requested section IDs, each with two or three paragraphs "
                                "of one hundred to eighteen hundred characters. Include no digits, URLs or citation IDs.")
                    if isinstance(exc, ValueError) and not isinstance(exc, ValidationError):
                        feedback = str(exc)
                    if attempt:
                        raise IntelligenceError(
                            "INVALID_REPORT_NARRATIVE",
                            "The writing model could not produce valid report content. Try generating the report again.",
                            retryable=True, details={"model_id": self.author.model_id, "reason": feedback},
                        ) from exc
                    payload["validation_feedback"] = feedback

        for section in content["sections"]:
            section_id = section["section_id"]
            if section_id not in written:
                continue
            original = section["paragraphs"]
            paragraphs = written[section_id]
            if section_id == "executive-summary":
                section["paragraphs"] = [*paragraphs, "Reviewed research snapshot: " + original[0]]
                section["claim_kinds"] = ["reviewed", "recommendation"]
            elif section_id in {"product-concept", "business-model", "market-scope"}:
                section["paragraphs"] = [original[0], *paragraphs]
                section["claim_kinds"] = list(dict.fromkeys([*section["claim_kinds"], "recommendation"]))
            else:
                section["paragraphs"] = paragraphs
                section["claim_kinds"] = ["recommendation"]
        return self.validate_content(content, context)

    def render(self, content: dict[str, Any], context: ProjectContext, format: str) -> dict[str, Any]:
        normalized = self.validate_content(content, context)
        slug = re.sub(r"[^a-zA-Z0-9]+", "-", normalized["title"]).strip("-").lower()[:60] or "startup-report"
        if format == "docx":
            data = _render_docx(normalized, context)
            return {"filename": f"{slug}.docx", "format": "docx",
                    "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "data": data}
        if format == "pdf":
            data = _render_pdf(normalized, context)
            return {"filename": f"{slug}.pdf", "format": "pdf", "media_type": "application/pdf", "data": data}
        raise IntelligenceError("UNSUPPORTED_FORMAT", "Report exports support DOCX and PDF.", 422,
                                details={"supported_formats": ["docx", "pdf"]})

    async def generate(self, kind: str, context: ProjectContext, options: dict[str, Any], progress) -> GenerationResult:
        if kind != "report":
            raise IntelligenceError("UNSUPPORTED_ARTIFACT_KIND", "The report module only generates reports.", 422)
        await progress("building_report", 20)
        content = self.build_content(context)
        if self.author is not None and self.author.enabled:
            content = await self._write_narrative(content, context, progress)
        await progress("rendering_report", 75)
        files = [self.render(content, context, fmt) for fmt in options["formats"]]
        artifact_id, file_ids = self.artifacts.create_artifact(context, "report", content, files)
        await progress("registering_report", 90)
        return GenerationResult(kind="report", context_ref={"project_id": context.project_id,
                                "context_version": context.context_version, "context_hash": context.content_hash},
                                artifact_ids=[artifact_id], file_ids=file_ids)


def _render_docx(content: dict[str, Any], context: ProjectContext) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72); section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82); section.right_margin = Inches(0.82)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor(37, 53, 59)
    normal.paragraph_format.space_after = Pt(7)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(content["title"]); run.bold = True; run.font.size = Pt(26); run.font.color.rgb = RGBColor(24, 51, 61)
    review_status = "ready" if context.critic.review_status == "ready" else "ready with gaps"
    meta = doc.add_paragraph(f"Prepared from reviewed project context v{context.context_version} · {review_status}")
    meta.runs[0].italic = True; meta.runs[0].font.color.rgb = RGBColor(27, 111, 102)
    source_index = {source.source_id: i + 1 for i, source in enumerate(context.sources)}
    for item in content["sections"]:
        doc.add_heading(item["heading"], level=1)
        if "recommendation" in item["claim_kinds"]:
            label = doc.add_paragraph("Proposed strategy and assumptions")
            label.runs[0].italic = True
        for paragraph in item["paragraphs"]:
            p = doc.add_paragraph(paragraph)
            cited = [source_index[sid] for sid in item["source_ids"] if sid in source_index]
            if cited:
                p.add_run(" " + "".join(f"[{index}]" for index in cited)).font.color.rgb = RGBColor(27, 111, 102)
        if item["evidence_ids"]:
            evidence_line = doc.add_paragraph("Evidence IDs: " + ", ".join(item["evidence_ids"]))
            evidence_line.runs[0].font.size = Pt(8)
            evidence_line.runs[0].font.color.rgb = RGBColor(90, 115, 112)
    if content["bibliography"]:
        doc.add_heading("Sources", level=1)
        for index, source in enumerate(content["bibliography"], start=1):
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.add_run(f"{source['title']} [{source['source_id']}]\n")
            if source.get("publisher"):
                paragraph.add_run(f"{source['publisher']} · ")
            paragraph.add_run(source["url"])
            paragraph.add_run(f"\nRetrieved: {source['retrieved_at']}")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"StartupForge · context v{context.context_version} · {context.content_hash[:12]}").font.size = Pt(8)
    output = io.BytesIO(); doc.save(output); return output.getvalue()


def _render_pdf(content: dict[str, Any], context: ProjectContext) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from xml.sax.saxutils import escape

    def register_unicode_font() -> tuple[str, str, str]:
        windows_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
        font_families = [
            (windows_fonts, "arial.ttf", "arialbd.ttf", "ariali.ttf"),
            (Path("/usr/share/fonts/truetype/liberation2"), "LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf", "LiberationSans-Italic.ttf"),
            (Path("/usr/share/fonts/truetype/dejavu"), "DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf"),
        ]
        for directory, regular, bold, italic in font_families:
            paths = [directory / regular, directory / bold, directory / italic]
            if not all(path.is_file() for path in paths):
                continue
            regular_name, bold_name, italic_name = "SFReportSans", "SFReportSans-Bold", "SFReportSans-Italic"
            pdfmetrics.registerFont(TTFont(regular_name, str(paths[0])))
            pdfmetrics.registerFont(TTFont(bold_name, str(paths[1])))
            pdfmetrics.registerFont(TTFont(italic_name, str(paths[2])))
            pdfmetrics.registerFontFamily(regular_name, normal=regular_name, bold=bold_name,
                                          italic=italic_name, boldItalic=bold_name)
            return regular_name, bold_name, italic_name
        return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"

    regular_font, bold_font, italic_font = register_unicode_font()
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=19*mm, leftMargin=19*mm, topMargin=18*mm, bottomMargin=18*mm,
                            title=content["title"], author="StartupForge")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SFTitle", parent=styles["Title"], fontName=bold_font, fontSize=23,
                              leading=28, textColor=colors.HexColor("#18333d"), alignment=TA_LEFT, spaceAfter=8))
    styles.add(ParagraphStyle(name="SFHeading", parent=styles["Heading2"], fontName=bold_font, fontSize=14, leading=18,
                              textColor=colors.HexColor("#1b6f66"), spaceBefore=14, spaceAfter=6,
                              keepWithNext=True))
    styles.add(ParagraphStyle(name="SFBody", parent=styles["BodyText"], fontName=regular_font,
                              fontSize=9.5, leading=14, spaceAfter=7))
    styles.add(ParagraphStyle(name="SFItalic", parent=styles["Italic"], fontName=italic_font))
    story = [Paragraph(escape(content["title"]), styles["SFTitle"]),
             Paragraph(f"Context v{context.context_version} · Review status: {escape('ready' if context.critic.review_status == 'ready' else 'ready with gaps')}", styles["SFItalic"]),
             Spacer(1, 7*mm)]
    for item in content["sections"]:
        story.append(Paragraph(escape(item["heading"]), styles["SFHeading"]))
        if "recommendation" in item["claim_kinds"]:
            story.append(Paragraph("Proposed strategy and assumptions", styles["SFItalic"]))
        for paragraph in item["paragraphs"]:
            story.append(Paragraph(escape(paragraph).replace("\n", "<br/>"), styles["SFBody"]))
        if item["evidence_ids"]:
            story.append(Paragraph("Evidence IDs: " + escape(", ".join(item["evidence_ids"])), styles["SFItalic"]))
    if content["bibliography"]:
        story.append(PageBreak())
        story.append(Paragraph("Sources", styles["SFHeading"]))
        for index, source in enumerate(content["bibliography"], start=1):
            text = f"{index}. {source['title']} [{source['source_id']}]<br/>{source.get('publisher') or ''} {source['url']}<br/>Retrieved: {source['retrieved_at']}"
            story.append(Paragraph(escape(text).replace("&lt;br/&gt;", "<br/>"), styles["SFBody"]))
    def footer(canvas, _doc):
        canvas.saveState(); canvas.setFont(regular_font, 8); canvas.setFillColor(colors.HexColor("#52696d"))
        canvas.drawString(19*mm, 10*mm, f"StartupForge · context v{context.context_version} · {context.content_hash[:12]}")
        canvas.drawRightString(A4[0]-19*mm, 10*mm, str(canvas.getPageNumber())); canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
