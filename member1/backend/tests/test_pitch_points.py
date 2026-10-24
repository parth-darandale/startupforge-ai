"""Regression checks for fuller, concise pitches across project domains."""

import io
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation

IMPLEMENTATION_ROOT = Path(__file__).resolve().parents[3]
for module_backend in (IMPLEMENTATION_ROOT / "member2" / "backend", IMPLEMENTATION_ROOT / "member3" / "backend"):
    sys.path.insert(0, str(module_backend))

from app.modules.intelligence.contracts import AnalysisReport, AnalysisSection, Claim, ProjectContext
from app.modules.intelligence.errors import IntelligenceError
from startupforge_member2.visual_studio import VisualStudioAdapter


FIXTURE = IMPLEMENTATION_ROOT / "integration" / "fixtures" / "project_context.demo.json"
OPTIONS = {"template_id": "investor-brief", "slide_count": 10, "tone": "professional"}


def context_for_repair_service() -> ProjectContext:
    context = ProjectContext.model_validate(json.loads(FIXTURE.read_text(encoding="utf-8")))
    context.idea.project_name = "Repair Relay"
    context.idea.original_text = "Repair Relay helps apartment residents request repairs and track maintenance updates."
    context.idea.problem = "Apartment residents lose track of maintenance requests and appointment updates."
    context.idea.solution = "A shared repair request log connects residents with building maintenance teams."
    context.idea.target_users = ["apartment residents", "building maintenance teams"]
    context.idea.domain = "Property maintenance"
    context.idea.assumptions = ["Residents will use a shared request log."]
    context.idea.constraints = ["Maintenance teams need a simple mobile workflow."]
    return context


def adapter() -> VisualStudioAdapter:
    # These checks exercise real pure content/renderer behavior without persistence.
    return VisualStudioAdapter.__new__(VisualStudioAdapter)


def deck_text(deck: Presentation) -> str:
    return "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)


def test_non_food_pitch_preserves_founder_input_and_adds_distinct_proposed_points():
    context = context_for_repair_service()
    content = adapter().build_deck_content(context, OPTIONS)
    recorded_text = "\n".join(
        [slide["speaker_notes"] for slide in content["slides"]]
        + [block["text"] for slide in content["slides"] for block in slide["body_blocks"]]
    )

    for founder_input in (context.idea.problem, context.idea.solution, *context.idea.target_users):
        assert founder_input in recorded_text
    recommendations = [block["text"] for slide in content["slides"] for block in slide["body_blocks"]
                       if block["claim_kind"] == "recommendation"]
    assert len(set(recommendations)) >= 15
    assert sum(len(slide["body_blocks"]) >= 5 for slide in content["slides"][1:-1]) >= 5
    assert "campus" not in " ".join(recommendations).lower()
    assert "meal" not in " ".join(recommendations).lower()


def test_recommendation_label_still_rejects_an_unsourced_market_number():
    context = context_for_repair_service()
    content = adapter().build_deck_content(context, OPTIONS)
    content["slides"][1]["body_blocks"] = [{
        "type": "text", "text": "The market contains 50 million paying customers.",
        "source_ids": [], "evidence_ids": [], "claim_kind": "recommendation",
    }]

    with pytest.raises(IntelligenceError) as exc:
        adapter().validate_content("pitch_deck", content, context)

    assert exc.value.code == "UNSOURCED_NUMERIC_CLAIM"


def test_pitch_accepts_a_valid_context_with_many_founder_assumptions():
    context = context_for_repair_service()
    context.idea.assumptions = [
        f"Founder assumption {index}: "
        + "Test how residents and building teams will use the shared repair workflow. " * 3
        for index in range(30)
    ]
    content = adapter().build_deck_content(context, OPTIONS)

    assert len(content["slides"]) == 10
    assert any(context.idea.assumptions[0] == block["text"]
               for slide in content["slides"] for block in slide["body_blocks"])


def test_pptx_renders_all_twelve_body_blocks_instead_of_silently_truncating():
    context = context_for_repair_service()
    points = [f"{name} point: a distinct validation question." for name in (
        "Alpha", "Bravo", "Charlie", "Delta", "Echo", "Foxtrot",
        "Golf", "Hotel", "India", "Juliet", "Kilo", "Lima",
    )]
    content = {
        "title": context.idea.project_name, "tone": "professional",
        "slides": [
            {"slide_id": "cover", "title": context.idea.project_name,
             "body_blocks": [{"text": context.idea.solution, "claim_kind": "provided"}]},
            {"slide_id": "questions", "title": "Validation questions",
             "body_blocks": [{"text": point, "claim_kind": "recommendation"} for point in points]},
        ],
    }
    rendered = adapter().render("pitch_deck", content, context, "pptx")
    deck = Presentation(io.BytesIO(rendered["data"]))
    text = deck_text(deck)

    for point in points:
        assert point in text
    for shape in deck.slides[1].shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= deck.slide_width
        assert shape.top + shape.height <= deck.slide_height


def test_pptx_preserves_speaker_notes_and_clickable_source_links():
    context = context_for_repair_service()
    source = context.sources[0]
    note = "Keep the full source qualification in these speaker notes."
    content = {
        "title": context.idea.project_name, "tone": "professional",
        "slides": [
            {"slide_id": "cover", "title": context.idea.project_name,
             "body_blocks": [{"text": context.idea.solution, "claim_kind": "provided"}]},
            {"slide_id": "source", "title": "Evidence and source",
             "speaker_notes": note, "source_ids": [source.source_id],
             "body_blocks": [{"text": f"{source.title}\n{source.url}", "claim_kind": "provided",
                              "source_ids": [source.source_id]}]},
        ],
    }
    rendered = adapter().render("pitch_deck", content, context, "pptx")
    deck = Presentation(io.BytesIO(rendered["data"]))

    assert note in deck.slides[1].notes_slide.notes_text_frame.text
    assert any(run.hyperlink.address == source.url
               for shape in deck.slides[1].shapes if shape.has_text_frame
               for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)


def test_existing_pilot_evidence_is_retained_without_a_blanket_absence_claim():
    context = context_for_repair_service()
    evidence = context.evidence[0]
    pilot_result = "The recorded pilot included completed repair requests, repeat use, and measured delivery effort."
    evidence.excerpt = pilot_result
    context.figures = []
    context.final_analysis = AnalysisReport(
        summary="Recorded pilot results inform the next test.",
        sections=[AnalysisSection(section_id="pilot", heading="Recorded pilot evidence", claims=[
            Claim(claim_id="recorded-pilot", text=pilot_result, kind="supported",
                  evidence_ids=[evidence.evidence_id]),
            Claim(claim_id="pilot-expansion", text="Expansion demand remains unverified.", kind="unknown"),
            Claim(claim_id="pilot-assumption", text="Building teams may adopt the shared log.", kind="assumption"),
        ])],
        gaps=[],
    )
    context.critic.unresolved_gaps = []
    content = adapter().build_deck_content(context, OPTIONS)
    blocks = [block for slide in content["slides"] for block in slide["body_blocks"]]
    text = " ".join([block["text"] for block in blocks] + [slide["speaker_notes"] for slide in content["slides"]])
    result = next(block for block in blocks if block.get("claim_id") == "recorded-pilot")

    assert result["text"] == pilot_result
    assert result["claim_kind"] == "supported"
    assert result["evidence_ids"] == [evidence.evidence_id]
    assert evidence.source_id in result["source_ids"]
    evidence_slide = next(slide for slide in content["slides"] if slide["title"] == "What the evidence says")
    evidence_claims = [block for block in evidence_slide["body_blocks"] if block["type"] == "claim"]
    assert evidence_claims and all(block["claim_kind"] == "supported" for block in evidence_claims)
    assert pilot_result in " ".join(block["text"] for block in evidence_claims)
    assert "contains no baseline or success targets" not in text
    assert "No validated demand, pricing or operating baseline" not in text
    assert "no pilot evidence" not in text.lower()
