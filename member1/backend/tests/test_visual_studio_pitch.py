"""Focused checks for the real Member 2 pitch deck content and PPTX renderer."""

import io
import json
import sys
from pathlib import Path

from pptx import Presentation

IMPLEMENTATION_ROOT = Path(__file__).resolve().parents[3]
for module_backend in (IMPLEMENTATION_ROOT / "member2" / "backend", IMPLEMENTATION_ROOT / "member3" / "backend"):
    sys.path.insert(0, str(module_backend))

from app.modules.intelligence.contracts import ProjectContext
from startupforge_member2.visual_studio import VisualStudioAdapter


FIXTURE = IMPLEMENTATION_ROOT / "integration" / "fixtures" / "project_context.demo.json"


def demo_context() -> ProjectContext:
    return ProjectContext.model_validate(json.loads(FIXTURE.read_text(encoding="utf-8")))


def adapter_without_storage() -> VisualStudioAdapter:
    # Content building and rendering are pure with respect to artifact persistence.
    return VisualStudioAdapter.__new__(VisualStudioAdapter)


def test_default_pitch_has_ten_substantive_slides_and_a_source_list():
    context = demo_context()
    adapter = adapter_without_storage()

    content = adapter.build_deck_content(context, {
        "template_id": "investor-brief", "slide_count": 10, "tone": "professional",
    })
    rendered = adapter.render("pitch_deck", content, context, "pptx")
    deck = Presentation(io.BytesIO(rendered["data"]))

    assert len(content["slides"]) == len(deck.slides) == 10
    titles = [slide["title"] for slide in content["slides"]]
    assert titles[0] == context.idea.project_name
    assert "Market opportunity to validate" in titles
    assert "Next step and sources" == titles[-1]
    assert not any(title in {"Review note", "Additional findings"} for title in titles)
    assert all(source.source_id in content["slides"][-1]["source_ids"] for source in context.sources[:11])

    cover_text = [shape.text for shape in deck.slides[0].shapes if shape.has_text_frame]
    assert any(context.idea.project_name in text for text in cover_text)
    cover_title = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame and context.idea.project_name in shape.text)
    assert cover_title.text_frame.paragraphs[0].font.size.pt >= 36
    assert cover_title.text_frame.paragraphs[0].font.name == "Arial"
    assert any(run.hyperlink.address for shape in deck.slides[-1].shapes if shape.has_text_frame
               for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)
    body_sizes = [run.font.size.pt for shape in deck.slides[1].shapes if shape.has_text_frame
                  for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size]
    assert max(body_sizes) >= 20


def test_pitch_slide_count_remains_selectable_without_filler():
    context = demo_context()
    adapter = adapter_without_storage()

    for requested in (5, 11, 12):
        content = adapter.build_deck_content(context, {
            "template_id": "investor-brief", "slide_count": requested, "tone": "professional",
        })
        assert len(content["slides"]) == requested
        assert content["slides"][-1]["title"] == "Next step and sources"
        assert all("No additional reviewed finding is available" not in block["text"]
                   for slide in content["slides"] for block in slide["body_blocks"])
