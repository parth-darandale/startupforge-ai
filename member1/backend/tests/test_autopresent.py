"""Focused checks for the symbolic AutoPresent program integration."""

import copy
import sys
import textwrap
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest

IMPLEMENTATION_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(IMPLEMENTATION_ROOT / "member2" / "backend"))

from app.modules.intelligence.errors import IntelligenceError
import startupforge_member2.autopresent as autopresent_module
from startupforge_member2.autopresent import (
    AutoPresentClient, _ProgramError, _SlideProgram, _prompt, validate_elements,
)


PROGRAM = """
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(2), Inches(11.9), Inches(2))
card.fill.solid()
card.fill.fore_color.rgb = RGBColor(240, 248, 248)
title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.9), Inches(1))
title.text = '__TITLE__'
p = title.text_frame.paragraphs[0]
p.font.size = Pt(38)
p.font.bold = True
p.font.color.rgb = RGBColor(24, 38, 64)
body = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.6))
body.text_frame.word_wrap = True
p = body.text_frame.paragraphs[0]
p.text = '__BLOCK_0__'
p.font.size = Pt(24)
p.font.color.rgb = RGBColor(24, 38, 64)
p.alignment = PP_ALIGN.LEFT
rule = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.7), Inches(6.3), Inches(12.6), Inches(6.3))
rule.line.color.rgb = RGBColor(8, 127, 140)
vertical = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(12.7), Inches(2), Inches(12.7), Inches(4))
prs.save('ignored-output.pptx')
"""


def _parse(program=PROGRAM):
    return _SlideProgram({"TITLE": "Campus meals", "BLOCK_0": "A customer pilot to validate demand."}).parse(program)


_ORIGINAL_IMPORT = __import__


def _deny_import(*args, **kwargs):
    # ast.walk imports its own standard-library deque lazily. The slide's pptx,
    # os, or other generated import statements must never reach this function.
    if args[0] == "collections":
        return _ORIGINAL_IMPORT(*args, **kwargs)
    raise AssertionError("The symbolic parser must not execute generated Python imports.")


def test_representative_python_pptx_program_becomes_editable_primitives_without_io():
    with patch("builtins.open") as opened, patch("builtins.__import__", new=_deny_import):
        result = _parse()
    opened.assert_not_called()
    assert {item["key"] for item in result if item["type"] == "text"} == {"TITLE", "BLOCK_0"}
    assert next(item for item in result if item.get("key") == "TITLE")["font_size"] == 38
    lines = [item for item in result if item.get("shape") == "line"]
    assert len(lines) == 2
    assert any(item["width"] == 0 for item in lines)
    assert any(item["height"] == 0 for item in lines)


@pytest.mark.parametrize("malicious", [
    "import os\nos.remove('anything')",
    "from os import system\nsystem('anything')",
    "open('escape.txt', 'w')",
    "for value in [1, 2]:\n    open('escape.txt', 'w')",
    "def generate():\n    open('escape.txt', 'w')",
    "from pptx import Presentation\nprs=Presentation()\nprs.__class__",
])
def test_unsupported_python_is_rejected_without_loading_modules_or_opening_files(malicious):
    with patch("builtins.open") as opened, patch("builtins.__import__", new=_deny_import):
        with pytest.raises(_ProgramError):
            _parse(malicious)
    opened.assert_not_called()


def test_simple_main_wrapper_is_interpreted_without_executing_python():
    wrapped = ("def create_slide():\n" + textwrap.indent(PROGRAM.strip(), "    ") +
               "\n    return prs\n\nif __name__ == '__main__':\n    create_slide()\n")
    with patch("builtins.open") as opened, patch("builtins.__import__", new=_deny_import):
        assert _parse(wrapped) == _parse()
    opened.assert_not_called()


def test_helpers_accept_fixed_parameters_literal_defaults_and_symbolic_return_values():
    program = """
from pptx import Presentation
from pptx.util import Inches, Pt
def add_text(slide, text, y, size=24, *, bold=False):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(y), Inches(11.9), Inches(1.5))
    box.text = text
    paragraph = box.text_frame.paragraphs[0]
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    return box
def create_slide():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, '__TITLE__', 0.5, 38, bold=True)
    add_text(slide, '__BLOCK_0__', 2.2)
    return prs
if __name__ == '__main__':
    deck = create_slide()
    deck.save('ignored-output.pptx')
"""
    with patch("builtins.open") as opened, patch("builtins.__import__", new=_deny_import):
        result = _parse(program)
    opened.assert_not_called()
    assert {item["key"] for item in result} == {"TITLE", "BLOCK_0"}
    assert next(item for item in result if item["key"] == "TITLE")["bold"]
    assert next(item for item in result if item["key"] == "BLOCK_0")["font_size"] == 24


@pytest.mark.parametrize("unsafe,reason", [
    ("def unused():\n    open('escape.txt', 'w')\n", "documented drawing calls"),
    ("def unused():\n    import os\n", "Import"),
    ("def unused():\n    from os import system\n", "documented pptx imports"),
    ("def unused():\n    value.__class__\n", "Private"),
    ("def unused():\n    def nested():\n        pass\n", "FunctionDef"),
    ("def unused(*args):\n    pass\n", "fixed parameters"),
    ("@unknown\ndef unused():\n    pass\n", "decorators"),
    ("def unused(value=Presentation()):\n    pass\n", "literal"),
    ("def unused():\n    for value in [1]:\n        pass\n", "For"),
    ("def unused():\n    if True:\n        pass\n", "If"),
    ("if True:\n    pass\n", "If"),
])
def test_wrapper_support_keeps_unsafe_or_unbounded_constructs_rejected(unsafe, reason):
    with patch("builtins.open") as opened, patch("builtins.__import__", new=_deny_import):
        with pytest.raises(_ProgramError, match=reason):
            _parse(PROGRAM + "\n" + unsafe)
    opened.assert_not_called()


def test_recursive_helpers_fail_with_a_bounded_reason():
    program = """
def first():
    second()
def second():
    first()
first()
"""
    with pytest.raises(_ProgramError, match="cannot recurse"):
        _parse(program)


SLIDESLIB_PROGRAM = """
def generate_presentation(output_pptx):
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    add_title(slide, 'Campus Meals', font_size=18, font_color=(0, 0, 0), background_color=None)
    add_text(slide, '\\nPartner with local kitchens. \\n\\nFor College students',
             [0.9166666666666666, 1.9965277777777777, 11.5, 4.758681102362205],
             font_size=18, color=(0, 0, 0), background_color=None)
    prs.save(output_pptx)
"""


def _parse_library(program=SLIDESLIB_PROGRAM):
    return _SlideProgram({"TITLE": "Campus Meals", "BLOCK_0": "Partner with local kitchens.",
                          "BLOCK_1": "For College students"}).parse(program)


def test_authentic_slideslib_entrypoint_maps_exact_canonical_group_and_emu_canvas_without_io():
    with patch("builtins.open") as opened, patch("builtins.__import__", new=_deny_import):
        result = _parse_library()
    opened.assert_not_called()
    assert {item["key"] for item in result} == {"TITLE", "BLOCK_0+BLOCK_1"}
    assert next(item for item in result if item["key"] == "TITLE")["font_size"] == 28
    body = next(item for item in result if item["key"] != "TITLE")
    assert body["y"] + body["height"] == 6.65
    assert body["font_size"] == 18


def test_slideslib_combined_placeholders_preserve_all_content_keys():
    program = SLIDESLIB_PROGRAM.replace("'Campus Meals'", "'__TITLE__'").replace(
        "'\\nPartner with local kitchens. \\n\\nFor College students'", "'__BLOCK_0__\\n\\n__BLOCK_1__'")
    assert {item["key"] for item in _parse_library(program)} == {"TITLE", "BLOCK_0+BLOCK_1"}


@pytest.mark.parametrize("replacement,reason", [
    ("'__BLOCK_0__'", "Missing required"),
])
def test_slideslib_groups_reject_added_omitted_or_duplicate_source_content(replacement, reason):
    program = SLIDESLIB_PROGRAM.replace("'\\nPartner with local kitchens. \\n\\nFor College students'", replacement)
    with pytest.raises(_ProgramError, match=reason):
        _parse_library(program)


def test_slideslib_extra_capabilities_are_not_implicitly_enabled():
    for helper in ("add_image", "get_answer", "generate_image", "google_search_screenshot", "search_image"):
        program = SLIDESLIB_PROGRAM.replace("prs.save(output_pptx)", f"{helper}(slide, 'anything')")
        with pytest.raises(_ProgramError, match="documented drawing calls"):
            _parse_library(program)


def test_grouped_primitives_still_need_unique_complete_block_coverage():
    result = _parse_library()
    assert validate_elements(result, 2) == result
    invalid = copy.deepcopy(result)
    invalid[-1]["key"] = "BLOCK_0+BLOCK_0"
    with pytest.raises(IntelligenceError):
        validate_elements(invalid, 2)
    invalid[-1]["key"] = "BLOCK_0+BLOCK_99"
    with pytest.raises(IntelligenceError):
        validate_elements(invalid, 2)


COMPACT_SLIDESLIB_PROGRAM = SLIDESLIB_PROGRAM.replace("12192000", "9144000").replace(
    "6858000", "5143500").replace("11.5, 4.758681102362205", "8.25, 3.9645833333333335").replace(
    "    prs.save(output_pptx)", """    add_text(slide, '11', [4.69996719160105, 5.285405730533683, 0.6000656167979003, 0.3395669291338583], font_size=18)
    prs.save(output_pptx)""")


def test_authentic_compact_canvas_is_normalized_and_only_footer_page_counter_is_ignored():
    result = _parse_library(COMPACT_SLIDESLIB_PROGRAM)
    assert len(result) == 2
    body = next(item for item in result if item["key"] != "TITLE")
    assert body["width"] == pytest.approx(8.25 * 13.333333 / 10)
    assert body["y"] + body["height"] == pytest.approx(6.65)
    assert body["y"] == pytest.approx(1.9965277777777777 * 6.65 / 5.625)


@pytest.mark.parametrize("program", [
    COMPACT_SLIDESLIB_PROGRAM.replace("'11'", "'$250,000'"),
    COMPACT_SLIDESLIB_PROGRAM.replace("'11'", "'1000'"),
    COMPACT_SLIDESLIB_PROGRAM.replace("'11'", "'The market is 11'"),
    COMPACT_SLIDESLIB_PROGRAM.replace("5.285405730533683", "2.0"),
    COMPACT_SLIDESLIB_PROGRAM.replace("'\\nPartner with local kitchens. \\n\\nFor College students'", "'We need $250,000 to expand.'"),
])
def test_generated_claim_wording_is_discarded_and_only_immutable_source_keys_survive(program):
    result = _parse_library(program)
    assert {item["key"] for item in result} == {"TITLE", "BLOCK_0+BLOCK_1"}
    assert all("text" not in item and "canonical_text" not in item for item in result)
    assert "$250,000" not in str(result)


def test_native_generated_wording_can_supply_geometry_but_never_rendered_text():
    result = _parse(PROGRAM.replace("p.text = '__BLOCK_0__'", "p.text = 'Invented $250,000 market claim'"))
    assert {item["key"] for item in result if item["type"] == "text"} == {"TITLE", "BLOCK_0"}
    assert "250,000" not in str(result)


def test_title_only_design_does_not_invent_body_geometry():
    start = COMPACT_SLIDESLIB_PROGRAM.index("    add_text(slide, '\\nPartner")
    finish = COMPACT_SLIDESLIB_PROGRAM.index("    add_text(slide, '11'", start)
    program = COMPACT_SLIDESLIB_PROGRAM[:start] + COMPACT_SLIDESLIB_PROGRAM[finish:]
    with pytest.raises(_ProgramError, match="Missing required"):
        _parse_library(program)


def test_canonical_numeric_content_is_preserved_even_if_placed_in_the_footer_area():
    program = COMPACT_SLIDESLIB_PROGRAM.replace("8.25, 3.9645833333333335", "8.25, 2.5")
    result = _SlideProgram({"TITLE": "Campus Meals", "BLOCK_0": "Partner with local kitchens.",
                            "BLOCK_1": "For College students", "BLOCK_2": "11"}).parse(program)
    assert len(result) == 3
    assert any(item["key"] == "BLOCK_2" for item in result)


def test_prompt_includes_exact_canonical_text_and_bounds_large_unicode_inputs_with_placeholders():
    prompt = _prompt({}, {"TITLE": "Campus meals", "BLOCK_0": "A pilot."}, None)
    assert 'Body paragraph (__BLOCK_0__, 8 characters): "A pilot."' in prompt
    assert "Do not rewrite" in prompt
    assert len(prompt.encode("utf-8")) < 2000
    assert prompt.index('"Campus meals"') < prompt.index("Presentation()")
    assert prompt.index('"A pilot."') < prompt.index("Presentation()")
    assert "Example program" not in prompt and "Supported Python:" not in prompt
    assert "BODY textboxes" in prompt
    large = _prompt({}, {"TITLE": "Campus meals", **{f"BLOCK_{index}": "大" * 20_000 for index in range(12)}},
                    {"project": "x" * 50_000})
    assert len(large.encode("utf-8")) <= 9000
    for index in range(12):
        assert f"__BLOCK_{index}__" in large
    assert "full source text is retained by the app" in large


@pytest.mark.parametrize("program,reason", [
    (PROGRAM.replace("p.text = '__BLOCK_0__'", "p.text = '__TITLE__'"), "more than one"),
    (PROGRAM.replace("p.text = '__BLOCK_0__'", "p.text = ''"), "placeholder"),
    (PROGRAM.replace("Inches(2.2)", "Inches(6.2)"), "inside the canvas"),
    (PROGRAM.replace("p.font.color.rgb = RGBColor(24, 38, 64)", "p.font.color.rgb = RGBColor(245, 245, 245)"), "contrasting"),
    (PROGRAM.replace("Inches(2.2)", "Inches(0.7)"), "overlap"),
])
def test_bad_content_geometry_and_colors_fail_with_a_repairable_reason(program, reason):
    with pytest.raises(_ProgramError, match=reason):
        _parse(program)


def test_edited_primitives_accept_schema_defaults_and_cannot_drop_a_research_block():
    result = _parse()
    with_defaults = [{"key": None, "shape": "rect", "fill": None, "line_color": None,
                      "font_size": 22, "bold": False, "color": "182640", "align": "left", **item}
                     for item in result]
    assert validate_elements(with_defaults, 1) == result
    invalid = [item for item in copy.deepcopy(result) if item.get("key") != "BLOCK_0"]
    with pytest.raises(IntelligenceError) as error:
        validate_elements(invalid, 1)
    assert error.value.code == "AUTOPRESENT_INVALID_ELEMENTS"
    assert "BLOCK_0" in error.value.details["reason"]


def test_background_color_stops_above_the_application_footer():
    program = PROGRAM.replace("slide = prs.slides.add_slide(prs.slide_layouts[6])", """
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(248, 250, 250)
""")
    result = _parse(program)
    assert result[0]["type"] == "shape"
    assert result[0]["height"] == 6.65


@pytest.mark.asyncio
async def test_client_requests_actual_python_and_repairs_only_once(tmp_path, monkeypatch):
    diagnostics = tmp_path / "autopresent" / "last-invalid-program.txt"
    monkeypatch.setattr(autopresent_module, "_INVALID_PROGRAM_PATH", diagnostics)
    client = AutoPresentClient("http://localhost:9999/v1", "autopresent")
    client._complete = AsyncMock(side_effect=[("import os", False), (PROGRAM, False)])
    result = await client.design_slide({"title": "Campus meals", "body_blocks": [{"text": "A pilot."}]})
    assert result
    assert client._complete.await_count == 2
    first_prompt, second_prompt = [call.args[0] for call in client._complete.await_args_list]
    assert "generating Python programs" in first_prompt
    assert "__BLOCK_0__" in first_prompt
    assert "canonical_text" not in first_prompt
    assert "A pilot." in first_prompt
    assert "Repair your design" in second_prompt
    client._complete = AsyncMock(return_value=("import os", False))
    with pytest.raises(IntelligenceError) as error:
        await client.design_slide({"title": "Campus meals", "body_blocks": [{"text": "A pilot."}]})
    assert client._complete.await_count == 2
    assert error.value.code == "AUTOPRESENT_INVALID_PROGRAM"
    assert error.value.details["reason"]
    diagnostic_text = diagnostics.read_text(encoding="utf-8")
    assert "AutoPresent program rejected:" in diagnostic_text
    assert "import os" in diagnostic_text
    assert len(diagnostic_text) <= 48_000


@pytest.mark.asyncio
async def test_missing_body_repair_explicitly_requires_body_textboxes(tmp_path, monkeypatch):
    monkeypatch.setattr(autopresent_module, "_INVALID_PROGRAM_PATH", tmp_path / "last-invalid-program.txt")
    start = SLIDESLIB_PROGRAM.index("    add_text(")
    finish = SLIDESLIB_PROGRAM.index("    prs.save(output_pptx)", start)
    title_only = SLIDESLIB_PROGRAM[:start] + SLIDESLIB_PROGRAM[finish:]
    client = AutoPresentClient("http://localhost:9999/v1", "autopresent")
    client._complete = AsyncMock(side_effect=[(title_only, False), (PROGRAM, False)])
    await client.design_slide({"title": "Campus meals", "body_blocks": [{"text": "A pilot."}]})
    repair = client._complete.await_args_list[1].args[0]
    assert "readable BODY textboxes" in repair
    assert "Add add_text calls below the title" in repair


def test_invalid_program_diagnostics_are_bounded_and_cannot_choose_a_path(tmp_path, monkeypatch):
    diagnostics = tmp_path / "last-invalid-program.txt"
    monkeypatch.setattr(autopresent_module, "_INVALID_PROGRAM_PATH", diagnostics)
    autopresent_module._record_invalid_program("x" * 60_000, "why" * 1000)
    assert len(diagnostics.read_text(encoding="utf-8")) == 48_000
