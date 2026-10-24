"""Context-grounded pitch and poster generation, with editable structured output."""

from __future__ import annotations

import io
import re
from typing import Annotated, Any, Literal
from urllib.parse import urlsplit

from pydantic import BaseModel, ConfigDict, Field, ValidationError

from app.integration.contracts import GenerationResult, TemplateDescriptor
from app.integration.interfaces import VisualStudioAdapter as VisualStudioProtocol
from app.modules.intelligence.contracts import ProjectContext
from app.modules.intelligence.errors import IntelligenceError
from app.modules.intelligence.output_author import OutputAuthor
from startupforge_member3.core.artifacts import ArtifactStore


class ContentModel(BaseModel):
    model_config = ConfigDict(extra="forbid")


class BodyBlock(ContentModel):
    type: Literal["text", "claim", "figure", "gap"] = "text"
    text: str = Field(min_length=1, max_length=3000)
    source_ids: list[str] = Field(default_factory=list)
    evidence_ids: list[str] = Field(default_factory=list)
    claim_kind: Literal["provided", "supported", "assumption", "unknown", "warning", "recommendation"] | None = None
    claim_id: str | None = None
    figure_id: str | None = None


class GeneratedElement(ContentModel):
    type: Literal["text", "shape"]
    key: str | None = Field(default=None, max_length=200, pattern=r"^(TITLE|BLOCK_\d+(?:\+BLOCK_\d+)*)$")
    x: float = Field(ge=0, le=13.3334, allow_inf_nan=False)
    y: float = Field(ge=0, le=6.65, allow_inf_nan=False)
    width: float = Field(ge=0, le=13.3334, allow_inf_nan=False)
    height: float = Field(ge=0, le=6.65, allow_inf_nan=False)
    font_size: float = Field(default=22, ge=17, le=48, allow_inf_nan=False)
    bold: bool = False
    color: str = Field(default="182640", pattern=r"^[0-9A-Fa-f]{6}$")
    align: Literal["left", "center", "right"] = "left"
    shape: Literal["rect", "ellipse", "line"] = "rect"
    fill: str | None = Field(default=None, pattern=r"^[0-9A-Fa-f]{6}$")
    line_color: str | None = Field(default=None, pattern=r"^[0-9A-Fa-f]{6}$")


class DeckSlide(ContentModel):
    slide_id: str = Field(min_length=1, max_length=80)
    title: str = Field(min_length=1, max_length=200)
    body_blocks: list[BodyBlock] = Field(default_factory=list, max_length=12)
    speaker_notes: str = Field(default="", max_length=4000)
    source_ids: list[str] = Field(default_factory=list)
    evidence_ids: list[str] = Field(default_factory=list)
    layout: Literal["auto", "cover", "statement", "split", "process", "comparison", "scorecard", "evidence", "closing"] = "auto"
    visual_labels: list[str] = Field(default_factory=list, max_length=2)
    emphasis_index: int = Field(default=0, ge=0, le=11)
    generated_elements: list[GeneratedElement] | None = Field(default=None, max_length=80)


class DeckDesign(ContentModel):
    theme: Literal["midnight", "ocean", "forest", "ember"] = "midnight"
    generator: Literal["standard", "autopresent"] = "standard"
    model_id: str | None = Field(default=None, max_length=200)


class PitchDeckContent(ContentModel):
    title: str = Field(min_length=1, max_length=200)
    tone: Literal["professional", "bold", "academic"] = "professional"
    slides: list[DeckSlide] = Field(min_length=1, max_length=30)
    design: DeckDesign | None = None


class SlideDesignCopy(ContentModel):
    slide_id: str
    title: str = Field(min_length=1, max_length=100)
    layout: Literal["cover", "statement", "split", "process", "comparison", "scorecard", "evidence", "closing"]
    visual_labels: list[Annotated[str, Field(min_length=1, max_length=60)]] = Field(default_factory=list, max_length=2)
    points: list[Annotated[str, Field(min_length=1, max_length=220)]] = Field(min_length=2, max_length=5)
    speaker_notes: str = Field(min_length=1, max_length=1200)


class DeckDesignBatch(ContentModel):
    theme: Literal["midnight", "ocean", "forest", "ember"]
    slides: list[SlideDesignCopy] = Field(min_length=1, max_length=3)


class PosterSection(ContentModel):
    section_id: str = Field(min_length=1, max_length=80)
    heading: str = Field(min_length=1, max_length=200)
    body: str = Field(min_length=1, max_length=5000)
    source_ids: list[str] = Field(default_factory=list)
    evidence_ids: list[str] = Field(default_factory=list)
    claim_kind: Literal["provided", "supported", "assumption", "unknown", "warning"] | None = None
    figure_id: str | None = None


class PosterContent(ContentModel):
    title: str = Field(min_length=1, max_length=200)
    sections: list[PosterSection] = Field(min_length=1, max_length=20)


class VisualStudioAdapter(VisualStudioProtocol):
    def __init__(self, artifacts: ArtifactStore, author: OutputAuthor | None = None, autopresent_runtime=None):
        self.artifacts = artifacts
        self.author = author
        self.autopresent_runtime = autopresent_runtime

    def templates(self, kind: Literal["pitch_deck", "poster"] | None = None) -> list[TemplateDescriptor]:
        templates = [
            TemplateDescriptor(template_id="investor-brief", name="Investor pitch", kind="pitch_deck",
                               supported_formats=["pptx"],
                               supported_options={"slide_count": {"minimum": 5, "maximum": 12, "default": 10},
                                                  "tone": {"values": ["professional", "bold", "academic"], "default": "professional"}}),
            TemplateDescriptor(template_id="evidence-one-pager", name="Evidence one pager", kind="poster",
                               supported_formats=["txt"],
                               supported_options={"audience": {"maximum_characters": 500, "default": "Potential customers and reviewers"}}),
        ]
        if self.autopresent_runtime is not None and self.autopresent_runtime.installed:
            templates.insert(1, TemplateDescriptor(
                template_id="autopresent", name="AutoPresent · Llama 3.1 8B (experimental)", kind="pitch_deck",
                supported_formats=["pptx"],
                supported_options={"slide_count": {"minimum": 5, "maximum": 12, "default": 10},
                                   "tone": {"values": ["professional", "bold", "academic"], "default": "professional"}}))
        return [item for item in templates if kind is None or item.kind == kind]

    def validate_options(self, kind: str, options: dict[str, Any]) -> None:
        if kind == "pitch_deck":
            if options.get("template_id") not in {item.template_id for item in self.templates("pitch_deck")}:
                raise IntelligenceError("TEMPLATE_NOT_FOUND", "Choose an installed pitch deck template.", 422)
            if options.get("tone") not in {"professional", "bold", "academic"}:
                raise IntelligenceError("INVALID_OPTION", "Choose professional, bold or academic tone.", 422)
            if not 5 <= int(options.get("slide_count", 0)) <= 12:
                raise IntelligenceError("INVALID_OPTION", "Slide count must be from 5 to 12.", 422)
        elif kind == "poster":
            if options.get("template_id") != "evidence-one-pager":
                raise IntelligenceError("TEMPLATE_NOT_FOUND", "Choose an installed poster template.", 422)
            if not str(options.get("audience", "")).strip() or len(options["audience"]) > 500:
                raise IntelligenceError("INVALID_OPTION", "Poster audience must contain 1 to 500 characters.", 422)
        else:
            raise IntelligenceError("UNSUPPORTED_ARTIFACT_KIND", "This module only generates pitch decks and posters.", 422)

    def validate_content(self, kind: str, content: dict[str, Any], context: ProjectContext) -> dict[str, Any]:
        model = PitchDeckContent if kind == "pitch_deck" else PosterContent if kind == "poster" else None
        if model is None:
            raise IntelligenceError("UNSUPPORTED_ARTIFACT_KIND", "This module cannot edit that artifact type.", 422)
        try:
            normalized = model.model_validate(content).model_dump(mode="json")
        except ValidationError as exc:
            raise IntelligenceError("INVALID_ARTIFACT_CONTENT", "The edited content does not match this artifact type.", 422,
                                    details={"issues": exc.errors(include_input=False)}) from exc
        self._validate_references(kind, normalized, context)
        if kind == "pitch_deck":
            from .autopresent import validate_elements
            for slide in normalized["slides"]:
                if slide.get("generated_elements") is not None:
                    validate_elements(slide["generated_elements"], len(slide["body_blocks"]))
        return normalized

    def _validate_references(self, kind: str, content: dict[str, Any], context: ProjectContext) -> None:
        source_ids = {source.source_id for source in context.sources}
        evidence_by_id = {item.evidence_id: item for item in context.evidence}
        evidence_ids = set(evidence_by_id)
        for section in content.get("slides", content.get("sections", [])):
            refs_source = section.get("source_ids", [])
            refs_evidence = section.get("evidence_ids", [])
            for block in section.get("body_blocks", []):
                refs_source = [*refs_source, *block.get("source_ids", [])]
                refs_evidence = [*refs_evidence, *block.get("evidence_ids", [])]
                if block.get("claim_kind") == "supported" and not block.get("evidence_ids"):
                    raise IntelligenceError("UNSOURCED_SUPPORTED_CLAIM", "Supported claims must retain evidence references.", 422)
                if block.get("type") == "figure":
                    figure = next((item for item in context.figures if item.figure_id == block.get("figure_id")), None)
                    if figure is None or figure.display_value not in block.get("text", "") or figure.excerpt not in block.get("text", ""):
                        raise IntelligenceError("FIGURE_CONTEXT_MISMATCH", "Figure text must preserve the exact quoted value and excerpt.", 422)
                claim_id = block.get("claim_id")
                if claim_id:
                    claim = next((claim for analysis_section in context.final_analysis.sections
                                  for claim in analysis_section.claims if claim.claim_id == claim_id), None)
                    if claim is None or block.get("text") != claim.text or block.get("claim_kind") != claim.kind:
                        raise IntelligenceError("CLAIM_CONTEXT_MISMATCH", "Reviewed claims must remain tied to their original reviewed wording.", 422)
            if any(ref not in source_ids for ref in refs_source) or any(ref not in evidence_ids for ref in refs_evidence):
                raise IntelligenceError("INVALID_CONTENT_REFERENCE", "Output content references evidence outside the selected project context.", 422)
            if any(evidence_by_id[ref].source_id not in refs_source for ref in refs_evidence):
                raise IntelligenceError("INVALID_CONTENT_REFERENCE", "Every evidence reference must include its matching source reference.", 422)
        figure_by_id = {figure.figure_id: figure for figure in context.figures}
        if kind == "pitch_deck":
            blocks = [block for slide in content["slides"] for block in slide["body_blocks"]]
        else:
            blocks = content["sections"]
        # Keep statistic cards exact, and require ordinary numeric claims to retain a
        # selected-context citation. Founder-provided numbers remain explicitly labelled.
        for block in blocks:
            text = str(block.get("text", block.get("body", "")))
            block_kind = block.get("claim_kind")
            block_type = block.get("type")
            if block_type == "figure" or block.get("figure_id"):
                figure = figure_by_id.get(block.get("figure_id"))
                if figure is None or figure.display_value not in text or figure.excerpt not in text:
                    raise IntelligenceError("FIGURE_CONTEXT_MISMATCH", "Figure text must preserve the exact quoted value and excerpt.", 422)
                continue
            if block_kind in {"provided", "assumption", "unknown", "warning"}:
                continue
            references = block.get("evidence_ids", [])
            if re.search(r"(?<![A-Za-z])(?:\$|€|£|₹|INR\s*)?\d[\d,.]*(?:\s*%|\s*(?:million|billion|thousand|crore|lakh))?", text, re.I) and not references:
                raise IntelligenceError("UNSOURCED_NUMERIC_CLAIM", "Numeric claims need evidence from this context.", 422)

    def _claim_groups(self, context: ProjectContext) -> dict[str, list[dict[str, Any]]]:
        groups: dict[str, list[dict[str, Any]]] = {}
        for section in context.final_analysis.sections:
            for claim in section.claims:
                source_ids = sorted({
                    evidence.source_id for evidence in context.evidence if evidence.evidence_id in claim.evidence_ids
                })
                groups.setdefault(section.heading, []).append({
                    "text": claim.text, "kind": claim.kind, "claim_id": claim.claim_id,
                    "source_ids": source_ids, "evidence_ids": list(claim.evidence_ids),
                })
        return groups

    def build_deck_content(self, context: ProjectContext, options: dict[str, Any]) -> dict[str, Any]:
        idea = context.idea
        groupings = self._claim_groups(context)
        supported_blocks = []
        risk_blocks = []
        for heading, claims in groupings.items():
            target = risk_blocks if any(word in heading.lower() for word in ("risk", "gap", "assumption")) else supported_blocks
            for claim in claims:
                target.append(BodyBlock(type="claim", text=claim["text"], source_ids=claim["source_ids"],
                                        evidence_ids=claim["evidence_ids"], claim_kind=claim["kind"], claim_id=claim["claim_id"]))

        def dedupe_claims(blocks: list[BodyBlock]) -> list[BodyBlock]:
            compact: list[BodyBlock] = []
            positions: dict[tuple[str, str | None, str], int] = {}
            for block in blocks:
                key = (block.text, block.claim_kind, block.type)
                if key not in positions:
                    positions[key] = len(compact)
                    compact.append(block)
                    continue
                index = positions[key]
                current = compact[index]
                compact[index] = current.model_copy(update={
                    "source_ids": sorted(set(current.source_ids) | set(block.source_ids)),
                    "evidence_ids": sorted(set(current.evidence_ids) | set(block.evidence_ids)),
                })
            return compact

        supported_blocks = dedupe_claims(supported_blocks)
        risk_blocks = dedupe_claims(risk_blocks)
        supported_index = {(block.text, block.claim_kind, block.type): index
                           for index, block in enumerate(supported_blocks)}
        distinct_risks: list[BodyBlock] = []
        for block in risk_blocks:
            key = (block.text, block.claim_kind, block.type)
            if key not in supported_index:
                distinct_risks.append(block)
                continue
            index = supported_index[key]
            current = supported_blocks[index]
            supported_blocks[index] = current.model_copy(update={
                "source_ids": sorted(set(current.source_ids) | set(block.source_ids)),
                "evidence_ids": sorted(set(current.evidence_ids) | set(block.evidence_ids)),
            })
        risk_blocks = distinct_risks
        # Evidence slides contain findings, rather than repeated generic unknowns.
        supported_blocks = [block for block in supported_blocks if block.claim_kind == "supported"]
        figure_blocks = [BodyBlock(type="figure", text=f"{figure.display_value} — {figure.metric}\nQuoted evidence: {figure.excerpt}\nQualification: {figure.qualification}",
                                   source_ids=[figure.source_id], evidence_ids=[figure.evidence_id], claim_kind="supported",
                                   figure_id=figure.figure_id)
                         for figure in context.figures[:3]]
        gaps = list(dict.fromkeys([*context.final_analysis.gaps, *context.critic.unresolved_gaps]))
        source_order = {source.source_id: index + 1 for index, source in enumerate(context.sources)}
        slides: list[DeckSlide] = []

        def add(title: str, blocks: list[BodyBlock], note: str = "") -> None:
            blocks = blocks[:12]
            if len(note) > 4000:
                note = note[:3800].rsplit("\n", 1)[0] + "\nFurther detail remains in the saved project context."
            source_ids = sorted({source for block in blocks for source in block.source_ids},
                                key=lambda source_id: source_order.get(source_id, 10_000))
            evidence_ids = sorted({evidence for block in blocks for evidence in block.evidence_ids})
            slides.append(DeckSlide(slide_id=f"slide-{len(slides) + 1:02d}", title=title, body_blocks=blocks,
                                    speaker_notes=note, source_ids=source_ids, evidence_ids=evidence_ids))

        def plan(text: str) -> BodyBlock:
            return BodyBlock(type="text", text=text, claim_kind="recommendation")

        targets = ", ".join(idea.target_users) or "Target users are not specified in the saved idea."
        add(idea.project_name, [BodyBlock(type="text", text=idea.solution, claim_kind="provided"),
                                BodyBlock(type="text", text=f"For {targets} · {idea.domain}", claim_kind="provided")],
            "This deck follows the selected saved project context. Assumptions, sourced findings and open questions remain distinct.")
        add("The customer problem", [BodyBlock(type="text", text=idea.problem, claim_kind="provided"),
                                     BodyBlock(type="text", text=f"Intended users: {targets}", claim_kind="provided"),
                                     plan("Frequency: ask when the problem last occurred and how often it repeats."),
                                     plan("Current workaround: compare the options customers already use."),
                                     plan("Impact: record the time, cost or inconvenience customers describe."),
                                     plan("Priority: check which users have the strongest reason to change.")])
        solution_blocks = [BodyBlock(type="text", text=idea.solution, claim_kind="provided")]
        if idea.constraints:
            solution_blocks.append(BodyBlock(type="text", text="Starting constraints: " + "; ".join(idea.constraints[:3]), claim_kind="provided"))
        solution_blocks.extend([
            plan("Core task: test the smallest workflow that delivers the proposed benefit."),
            plan("User journey: check how a user starts, completes the task and gets the result."),
            plan("First version: prioritize essential features before adding optional features."),
            plan("Dependencies: confirm the partners, tools or resources needed for delivery."),
        ])
        add("The product concept", solution_blocks)

        evidence_blocks = supported_blocks[:2] + figure_blocks[:1]
        if not evidence_blocks:
            evidence_blocks = [BodyBlock(type="gap", text="The saved research contains no supported claim or sourced figure for this idea yet.",
                                         claim_kind="unknown")]
        evidence_blocks.extend([
            plan("Source scope: check the date, location and population behind each finding."),
            plan("Customer fit: test whether the cited finding applies to the intended users."),
            plan("Evidence gaps: collect direct observations for claims the sources cannot settle."),
        ])
        add("What the evidence says", evidence_blocks,
            "Supported claims and figures retain their source and evidence references. Read each figure qualification before reuse.")

        market_figure_ids = {figure.figure_id for figure in context.figures
                             if any(term in figure.metric.lower() for term in ("market", "enrollment", "spend", "revenue", "industry", "sales"))}
        market_figures = [figure for figure in figure_blocks if figure.figure_id in market_figure_ids]
        market_blocks = list(market_figures[:3])
        if not market_blocks:
            market_blocks.append(BodyBlock(type="gap", text="No market-size figure is separately recorded in this context.", claim_kind="unknown"))
        market_blocks.extend([
            plan("Eligible users: define the customer group and geographic boundary."),
            plan("Reachable segment: estimate how many eligible users the team can serve."),
            plan("Frequency: test how often customers would use or purchase the offer."),
            plan("Value: check willingness to pay where payment is part of the model."),
            plan("Alternatives: compare price, convenience and outcomes with existing options."),
        ])
        add("Market opportunity to validate", market_blocks)

        business_blocks = [BodyBlock(type="text", text=item, claim_kind="assumption")
                           for item in idea.assumptions[:1]]
        business_blocks.extend([
            plan("Payer and offer: identify who pays and what they receive, if applicable."),
            plan("Pricing: test the payment amount, frequency and willingness to commit."),
            plan("Acquisition: compare ways to reach users and the effort each channel needs."),
            plan("Delivery cost: record direct costs, partner fees and support effort."),
            plan("Sustainability: compare revenue with delivery costs and repeat use."),
        ])
        add("Business model hypothesis", business_blocks,
            "Founder assumptions in the saved context:\n" + "\n".join(idea.assumptions))

        add("A focused pilot", [
            plan(f"Participants: recruit a focused group from {targets}."),
            plan("Baseline: record the recent problem and the workaround each participant uses."),
            plan("Task test: observe participants completing the core workflow."),
            plan("Commitment: request a concrete next step appropriate to the business model."),
            plan("Delivery: log delays, direct costs and support needed during the test."),
            plan("Follow-up: check repeat use and use feedback to select the next change."),
        ], "These are proposed validation steps, not results or commitments already made.")

        add("Pilot scorecard", [
            plan("Activation: participants who complete the core task."),
            plan("Commitment: users who take the intended next step."),
            plan("Retention: participants who return or repeat the task."),
            plan("Outcome: improvement in the customer problem the offer targets."),
            plan("Reliability: completion failures, delays and support requests."),
            plan("Economics: direct cost and operating effort per completed use."),
        ], "Suggested measures for a pilot. Compare them with any existing baseline and set success criteria before the next test.")

        decision_blocks = [block for block in risk_blocks if not block.text.startswith("This conclusion remains unknown")][:2]
        if gaps:
            decision_blocks.extend(BodyBlock(type="gap", text=item, claim_kind="unknown") for item in gaps[:2])
        decision_blocks.extend([
            plan("Demand risk: compare demonstrated commitment with the pilot target."),
            plan("Delivery risk: confirm capacity and critical partner dependencies."),
            plan("Decision owner: assign responsibility for each unresolved launch condition."),
            plan("Expansion gate: define what result supports proceeding, revising or stopping."),
        ])
        decision_blocks = dedupe_claims(decision_blocks)
        add("Risks and decision gates", decision_blocks)

        reference_blocks = [plan("Next step: test the highest-risk assumption and record the decision it informs.")]
        for source in context.sources[:11]:
            source_number = source_order[source.source_id]
            reference_blocks.append(BodyBlock(type="text", text=f"[{source_number}] {source.title}\n{urlsplit(str(source.url)).netloc}",
                                              source_ids=[source.source_id], claim_kind="provided"))
        if not context.sources:
            reference_blocks.append(BodyBlock(type="gap", text="No source records are attached to this context. Add relevant primary sources before using market claims.", claim_kind="unknown"))
        add("Next step and sources", reference_blocks,
            f"Source numbers map to the saved context. Review status: {'ready' if context.critic.review_status == 'ready' else 'ready with gaps'}.")

        extras: list[DeckSlide] = []
        remaining_evidence = supported_blocks[2:]
        if remaining_evidence:
            extras.append(DeckSlide(slide_id="appendix-evidence", title="Additional reviewed findings",
                                    body_blocks=remaining_evidence[:6], speaker_notes="Additional claims copied from the reviewed analysis."))
        else:
            extras.append(DeckSlide(slide_id="appendix-evidence", title="Further evidence to gather",
                                    body_blocks=[BodyBlock(type="gap", text="Collect direct customer, competitor and operating evidence before making stronger claims.", claim_kind="unknown")],
                                    speaker_notes="No additional supported finding is available in this context."))
        extra_gaps = gaps[2:]
        extras.append(DeckSlide(slide_id="appendix-gaps", title="Open validation questions",
                                body_blocks=[BodyBlock(type="gap", text=gap, claim_kind="unknown") for gap in extra_gaps[:8]] or
                                             [BodyBlock(type="gap", text="No further open questions were recorded in the saved analysis.", claim_kind="unknown")],
                                speaker_notes="These questions remain unresolved in the saved context."))

        count = int(options["slide_count"])
        selected = [*slides[:-1][:count - 1], slides[-1]]
        if count > 10:
            selected = [*slides[:-1], *extras[:count - 10], slides[-1]]
        return self.validate_content("pitch_deck", {"title": idea.project_name, "tone": options["tone"],
                                                       "slides": [s.model_dump(mode="json") for s in selected]}, context)

    async def author_deck_content(self, content: dict[str, Any], context: ProjectContext,
                                  options: dict[str, Any], progress) -> dict[str, Any]:
        """The model designs the composition; research blocks remain immutable."""
        if self.author is None or not self.author.enabled:
            return content
        slides = content["slides"]
        outline = [{"slide_id": slide["slide_id"], "purpose": slide["title"]} for slide in slides]
        theme = None
        layouts: list[str] = []
        for start in range(0, len(slides), 3):
            batch = slides[start:start + 3]
            required_ids = [slide["slide_id"] for slide in batch]
            schema = DeckDesignBatch.model_json_schema()
            schema["$defs"]["SlideDesignCopy"]["properties"]["slide_id"] = {"type": "string", "enum": required_ids}
            schema["properties"]["slides"].update(minItems=len(batch), maxItems=len(batch))
            await progress(f"ai_designing_slides_{start + 1}_to_{start + len(batch)}",
                           25 + int(40 * start / len(slides)))
            payload = {
                "idea": context.idea.model_dump(mode="json"), "tone": options["tone"],
                "deck_outline": [item["purpose"] for item in outline],
                "required_slide_ids": required_ids, "selected_theme": theme, "previous_layouts": layouts,
                "open_questions": list(dict.fromkeys([*context.final_analysis.gaps,
                                                       *context.critic.unresolved_gaps])),
                "slide_briefs": [{"slide_id": slide["slide_id"], "purpose": slide["title"],
                                  "fixed_content": [block for block in slide["body_blocks"]
                                                    if block.get("claim_kind") not in {"recommendation", "supported"}
                                                    and not block.get("source_ids")],
                                  "retained_evidence_blocks": sum(bool(block.get("evidence_ids"))
                                                                   for block in slide["body_blocks"]),
                                  "writing_brief": [block["text"] for block in slide["body_blocks"]
                                                    if block.get("claim_kind") == "recommendation"]}
                                 for slide in batch],
            }
            draft = None
            for attempt in range(2):
                try:
                    result = await self.author.generate_json("pitch_deck_copy", payload,
                                                             schema)
                    draft = DeckDesignBatch.model_validate(result)
                    expected = {slide["slide_id"] for slide in batch}
                    if len(draft.slides) != len(batch) or {slide.slide_id for slide in draft.slides} != expected:
                        raise ValueError("Return every requested slide_id exactly once: " + ", ".join(required_ids))
                    def usable(text: str) -> bool:
                        return not re.search(r"(?<![A-Za-z])\d|https?://|\b(?:ev|src|fig)_[\w-]+",
                                             text)
                    baseline_by_id = {slide["slide_id"]: slide for slide in batch}
                    for item in draft.slides:
                        item.points = list(dict.fromkeys(point.strip() for point in item.points if usable(point)))
                        if len(item.points) < 2:
                            raise ValueError("Return at least two useful points without numbers, URLs or evidence IDs; exact figures are merged separately.")
                        if not usable(item.title):
                            item.title = baseline_by_id[item.slide_id]["title"]
                        item.visual_labels = [label for label in item.visual_labels if usable(label)]
                        if not usable(item.speaker_notes):
                            item.speaker_notes = ""
                    break
                except (ValidationError, ValueError) as exc:
                    if attempt:
                        reason = (str(exc) if not isinstance(exc, ValidationError) else
                                  str(exc.errors(include_input=False)))
                        raise IntelligenceError("INVALID_SLIDE_DESIGN", "The model could not produce a usable slide design. Please try again.",
                                                502, retryable=True, details={"reason": reason[:1200]}) from exc
                    payload["validation_feedback"] = str(exc)[:700]
                except IntelligenceError as exc:
                    if attempt or exc.code not in {"MODEL_INVALID_JSON", "MODEL_TRUNCATED"}:
                        raise
                    payload["validation_feedback"] = "Return a shorter complete JSON object matching the schema."
            if draft is None:
                raise IntelligenceError("INVALID_SLIDE_DESIGN", "No slide design was returned.", 502)
            theme = theme or draft.theme
            by_id = {item.slide_id: item for item in draft.slides}
            for offset, slide in enumerate(batch):
                item = by_id[slide["slide_id"]]
                position = start + offset
                fixed = [block for block in slide["body_blocks"] if block.get("claim_kind") != "recommendation"]
                if position == 0:
                    # The cover displays a model-authored value proposition and audience line.
                    fixed = []
                    points = item.points[:2]
                    slide["layout"] = "cover"
                else:
                    points = item.points[:max(0, 12 - len(fixed))]
                    slide["layout"] = ("closing" if position == len(slides) - 1 else
                                       "statement" if item.layout in {"cover", "closing"} else item.layout)
                slide["title"] = context.idea.project_name if position == 0 else item.title
                slide["body_blocks"] = [*fixed, *[BodyBlock(type="text", text=point.strip(), claim_kind="recommendation").model_dump(mode="json")
                                                 for point in points]]
                slide["visual_labels"] = item.visual_labels
                slide["emphasis_index"] = next((i for i, block in enumerate(slide["body_blocks"])
                                                if block["type"] == "figure"), 0)
                retained_note = slide.get("speaker_notes", "")
                extra_points = "\n".join(item.points[len(points):])
                generated_note = "\n".join(part for part in [item.speaker_notes, extra_points] if part)
                room = max(0, 3998 - len(retained_note))
                slide["speaker_notes"] = (retained_note + "\n\n" + generated_note[:room]).strip()
                layouts.append(slide["layout"])
        content["design"] = {"theme": theme}
        return self.validate_content("pitch_deck", content, context)

    def build_poster_content(self, context: ProjectContext, options: dict[str, Any]) -> dict[str, Any]:
        idea = context.idea
        sections: list[PosterSection] = [
            PosterSection(section_id="problem", heading="Problem", body=idea.problem, claim_kind="provided"),
            PosterSection(section_id="solution", heading="Solution", body=idea.solution, claim_kind="provided"),
            PosterSection(section_id="audience", heading="Audience", body=", ".join(idea.target_users) or "Target users are not specified.", claim_kind="provided"),
        ]
        for heading, claims in self._claim_groups(context).items():
            entries = []
            sources: set[str] = set()
            evidence: set[str] = set()
            kinds: set[str] = set()
            for claim in claims[:5]:
                label = {"supported": "Supported", "assumption": "Assumption", "unknown": "Unknown"}[claim["kind"]]
                entries.append(f"[{label}] {claim['text']}")
                sources.update(claim["source_ids"])
                evidence.update(claim["evidence_ids"])
                kinds.add(claim["kind"])
            if entries:
                sections.append(PosterSection(section_id=f"finding-{len(sections) + 1}", heading=heading,
                                              body="\n".join(entries), source_ids=sorted(sources), evidence_ids=sorted(evidence),
                                              claim_kind="supported" if "supported" in kinds else "assumption" if "assumption" in kinds else "unknown"))
        for figure in context.figures[:6]:
            sections.append(PosterSection(section_id=f"figure-{figure.figure_id}", heading=figure.metric[:120],
                                          body=f"{figure.display_value}\n{figure.excerpt}\nQualification: {figure.qualification}",
                                          source_ids=[figure.source_id], evidence_ids=[figure.evidence_id], claim_kind="supported",
                                          figure_id=figure.figure_id))
        gaps = list(dict.fromkeys([*context.final_analysis.gaps, *context.critic.unresolved_gaps]))
        if gaps:
            sections.append(PosterSection(section_id="gaps", heading="Open questions", body="\n".join(gaps[:8]), claim_kind="unknown"))
        sections.append(PosterSection(section_id="source-note", heading="Evidence and review", body=(
            f"Review status: {'ready' if context.critic.review_status == 'ready' else 'ready with gaps'}. Prepared for {options['audience']}. "
            "Verify cited sources before external publication."), claim_kind="provided"))
        return self.validate_content("poster", {"title": idea.project_name, "sections": [s.model_dump(mode="json") for s in sections]}, context)

    def render(self, kind: str, content: dict[str, Any], context: ProjectContext, format: str) -> dict[str, Any]:
        normalized = self.validate_content(kind, content, context)
        if kind == "pitch_deck" and format == "pptx":
            if normalized.get("design"):
                from .deck_design import render_designed_pptx
                data = render_designed_pptx(normalized, context)
            else:
                data = _render_pptx(normalized, context)
            return {"filename": f"{_slug(normalized['title'])}-pitch.pptx", "format": "pptx",
                    "media_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    "data": data}
        if kind == "poster" and format == "txt":
            return {"filename": f"{_slug(normalized['title'])}-poster.txt", "format": "txt",
                    "media_type": "text/plain; charset=utf-8", "data": _render_poster_txt(normalized, context)}
        raise IntelligenceError("UNSUPPORTED_FORMAT", "This visual studio template does not support that export format.", 422)

    async def generate(self, kind: str, context: ProjectContext, options: dict[str, Any], progress) -> GenerationResult:
        await progress("building_editable_content", 25)
        content = self.build_deck_content(context, options) if kind == "pitch_deck" else self.build_poster_content(context, options)
        if kind == "pitch_deck" and options.get("template_id") == "autopresent":
            await progress("loading_autopresent", 28)
            async with self.autopresent_runtime.session() as designer:
                content["design"] = {"theme": "midnight", "generator": "autopresent", "model_id": designer.model}
                for index, slide in enumerate(content["slides"]):
                    await progress(f"autopresent_slide_{index + 1}_of_{len(content['slides'])}",
                                   30 + int(48 * index / len(content["slides"])))
                    slide["generated_elements"] = await designer.design_slide(
                        slide, {"project": content["title"], "tone": options["tone"],
                                "slide_number": index + 1, "slide_count": len(content["slides"]),
                                "theme": "White background, navy text, teal accents. Consistent modern investor pitch."})
        elif kind == "pitch_deck":
            content = await self.author_deck_content(content, context, options, progress)
        await progress("rendering_export", 82)
        output = self.render(kind, content, context, "pptx" if kind == "pitch_deck" else "txt")
        artifact_id, file_ids = self.artifacts.create_artifact(context, kind, content, [output])
        await progress("registering_artifact", 90)
        return GenerationResult(kind=kind, context_ref={"project_id": context.project_id,
                                "context_version": context.context_version, "context_hash": context.content_hash},
                                artifact_ids=[artifact_id], file_ids=file_ids)


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", value).strip("-").lower()
    return slug[:60] or "startupforge"


def _render_poster_txt(content: dict[str, Any], context: ProjectContext) -> bytes:
    sources = {source.source_id: source for source in context.sources}
    lines = [content["title"], "=" * min(72, len(content["title"])), ""]
    for section in content["sections"]:
        lines.extend([section["heading"], "-" * min(48, len(section["heading"])), section["body"], ""])
        for source_id in section["source_ids"]:
            source = sources.get(source_id)
            if source:
                lines.append(f"Source: {source.title} — {source.url} (ID: {source_id})")
        if section["evidence_ids"]:
            lines.append("Evidence IDs: " + ", ".join(section["evidence_ids"]))
        if section["source_ids"]:
            lines.append("")
    review_status = "ready" if context.critic.review_status == "ready" else "ready with gaps"
    lines.append(f"Context version: {context.context_version} | Review: {review_status}")
    return "\n".join(lines).encode("utf-8")


def _render_pptx(content: dict[str, Any], context: ProjectContext) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]
    tone = content.get("tone", "professional")
    navy = RGBColor(31, 45, 59) if tone == "academic" else RGBColor(33, 40, 44) if tone == "bold" else RGBColor(17, 52, 49)
    teal = RGBColor(69, 112, 161) if tone == "academic" else RGBColor(220, 118, 65) if tone == "bold" else RGBColor(27, 132, 119)
    paper = RGBColor(247, 247, 241)
    ink = RGBColor(35, 49, 54)
    muted = RGBColor(83, 103, 105)
    light = RGBColor(235, 242, 237)
    white = RGBColor(255, 255, 255)
    label_color = {"supported": teal, "assumption": RGBColor(182, 112, 43),
                   "unknown": muted, "warning": RGBColor(176, 67, 58), "provided": navy}
    source_numbers = {source.source_id: index + 1 for index, source in enumerate(context.sources)}
    sources = {source.source_id: source for source in context.sources}
    deck.core_properties.title = content["title"]
    deck.core_properties.subject = "Startup validation pitch grounded in the selected project context"
    deck.core_properties.author = "StartupForge"

    for index, slide_data in enumerate(content["slides"]):
        slide = deck.slides.add_slide(blank)
        is_cover = index == 0
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = navy if is_cover else paper

        if is_cover:
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), Inches(7.5))
            accent.fill.solid(); accent.fill.fore_color.rgb = teal; accent.line.fill.background()
            eyebrow = slide.shapes.add_textbox(Inches(0.95), Inches(0.75), Inches(10.8), Inches(0.36)).text_frame
            eyebrow.text = "STARTUPFORGE  /  VALIDATION PITCH"
            eyebrow.paragraphs[0].font.name = "Arial"; eyebrow.paragraphs[0].font.size = Pt(13)
            eyebrow.paragraphs[0].font.bold = True; eyebrow.paragraphs[0].font.color.rgb = light
            title_box = slide.shapes.add_textbox(Inches(0.95), Inches(1.55), Inches(11.1), Inches(1.45)).text_frame
            title_box.clear(); title_box.word_wrap = True; title_box.vertical_anchor = MSO_ANCHOR.MIDDLE
            title_p = title_box.paragraphs[0]; title_p.text = slide_data["title"]
            title_p.font.name = "Arial"; title_p.font.size = Pt(42 if len(slide_data["title"]) < 30 else 36)
            title_p.font.bold = True; title_p.font.color.rgb = white
            subtitle = slide_data["body_blocks"][0]["text"] if slide_data["body_blocks"] else ""
            sub_box = slide.shapes.add_textbox(Inches(0.98), Inches(3.2), Inches(10.75), Inches(1.65)).text_frame
            sub_box.clear(); sub_box.word_wrap = True
            sub_p = sub_box.paragraphs[0]; sub_p.text = subtitle
            sub_p.font.name = "Arial"; sub_p.font.size = Pt(24); sub_p.font.color.rgb = light
            sub_p.space_after = Pt(0)
            if len(slide_data["body_blocks"]) > 1:
                line = slide.shapes.add_textbox(Inches(0.98), Inches(5.2), Inches(10.6), Inches(0.58)).text_frame
                line.text = slide_data["body_blocks"][1]["text"]
                line.paragraphs[0].font.name = "Arial"; line.paragraphs[0].font.size = Pt(17)
                line.paragraphs[0].font.color.rgb = light
            footer_color = light
        else:
            # A consistent title rule and aligned content area establish the deck theme.
            title_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.42), Inches(11.45), Inches(0.84)).text_frame
            title_box.clear(); title_box.word_wrap = True; title_box.vertical_anchor = MSO_ANCHOR.MIDDLE
            title_p = title_box.paragraphs[0]; title_p.text = slide_data["title"]
            title_p.font.name = "Arial"; title_p.font.size = Pt(31 if len(slide_data["title"]) < 42 else 28)
            title_p.font.bold = True; title_p.font.color.rgb = navy
            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.34), Inches(11.45), Inches(0.035))
            rule.fill.solid(); rule.fill.fore_color.rgb = teal; rule.line.fill.background()

            blocks = slide_data["body_blocks"]
            body_top, body_bottom = 1.72, 6.48
            available = body_bottom - body_top
            # Dense lists use columns so every saved point remains visible.
            column_count = 2 if len(blocks) > 6 else 1
            column_width = 5.35 if column_count == 2 else 11.05
            per_column = (len(blocks) + column_count - 1) // column_count
            placements = []
            for column in range(column_count):
                column_blocks = blocks[column * per_column:(column + 1) * per_column]
                desired = []
                for block in column_blocks:
                    text = block["text"]
                    width_ratio = column_width / 11.05
                    if block["type"] == "figure":
                        value, *details = text.split("\n")
                        chars = len(value) / (65 * width_ratio) + sum(len(item) / (104 * width_ratio) for item in details)
                        font_size = 22
                    else:
                        chars = len(text) / (92 * width_ratio)
                        font_size = 21 if len(text) < 150 else 19 if len(text) < 380 else 17
                    desired.append((max(0.58, chars * (font_size / 72) * 1.32 + 0.2), font_size))
                gap = 0.16 if len(column_blocks) <= 4 else 0.10
                total = sum(height for height, _ in desired) + max(0, len(column_blocks) - 1) * gap
                if total > available and desired:
                    space_for_text = available - max(0, len(column_blocks) - 1) * gap
                    factor = space_for_text / sum(height for height, _ in desired)
                    desired = [(height * factor, max(16, size - (2 if factor < 0.82 else 1))) for height, size in desired]
                    total = available
                y = body_top + max(0, (available - total) * 0.08)
                for height, font_size in desired:
                    placements.append((1.0 + column * 5.7, y, height, font_size))
                    y += height + gap
            for block, (x, y, height, font_size) in zip(blocks, placements):
                if block["type"] == "figure":
                    value, *details = block["text"].split("\n")
                    value_box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(column_width), Inches(min(0.58, height))).text_frame
                    value_box.word_wrap = True
                    value_p = value_box.paragraphs[0]; value_p.text = value
                    value_p.font.name = "Arial"; value_p.font.size = Pt(28); value_p.font.bold = True
                    value_p.font.color.rgb = teal
                    y += min(0.64, height)
                    detail_text = "\n".join(details)
                    if detail_text:
                        detail_box = slide.shapes.add_textbox(Inches(x + 0.02), Inches(y), Inches(column_width - 0.02), Inches(max(0.54, height - 0.5))).text_frame
                        detail_box.word_wrap = True
                        detail_p = detail_box.paragraphs[0]; detail_p.text = detail_text
                        detail_p.font.name = "Arial"; detail_p.font.size = Pt(16 if len(detail_text) > 260 else 17)
                        detail_p.font.color.rgb = ink
                        y += max(0.54, height - 0.5)
                else:
                    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(column_width), Inches(height)).text_frame
                    box.clear(); box.word_wrap = True; box.margin_left = Inches(0.02); box.margin_right = Inches(0.02)
                    box.margin_top = Inches(0.01); box.margin_bottom = Inches(0.02)
                    paragraph = box.paragraphs[0]
                    kind = block.get("claim_kind") or "provided"
                    label = {"supported": "SUPPORTED  ·  ", "assumption": "ASSUMPTION  ·  ",
                             "unknown": "OPEN  ·  ", "warning": "REVIEW  ·  ",
                             "recommendation": "PLAN  ·  "}.get(kind, "")
                    if label:
                        run = paragraph.add_run(); run.text = label
                        run.font.name = "Arial"; run.font.size = Pt(12); run.font.bold = True
                        run.font.color.rgb = label_color.get(kind, teal)
                    for line_index, line_text in enumerate(block["text"].split("\n")):
                        if line_index:
                            paragraph.add_line_break()
                        run = paragraph.add_run(); run.text = line_text
                        run.font.name = "Arial"; run.font.size = Pt(font_size)
                        run.font.color.rgb = muted if block["type"] == "gap" else ink
                        source_link = sources.get(block["source_ids"][0]) if block.get("source_ids") and block["text"].startswith("[") else None
                        if source_link or line_text.startswith(("https://", "http://")):
                            run.hyperlink.address = str(source_link.url) if source_link else line_text
                            run.font.color.rgb = teal
                            run.font.underline = True
                    paragraph.space_after = Pt(0); paragraph.line_spacing = 1.06

            source_ids = sorted(set(slide_data["source_ids"]),
                                key=lambda source_id: source_numbers.get(source_id, 10_000))
            evidence_ids = sorted(set(slide_data["evidence_ids"]))
            references = [f"[{source_numbers[source_id]}]" for source_id in source_ids if source_id in source_numbers]
            footer_text = "Sources " + ", ".join(references) if references else ""
            if evidence_ids:
                footer_text += ("  ·  " if footer_text else "") + "Evidence retained in project context"
            if footer_text:
                citation = slide.shapes.add_textbox(Inches(1.0), Inches(6.62), Inches(10.7), Inches(0.25)).text_frame
                citation.text = footer_text
                citation.paragraphs[0].font.name = "Arial"; citation.paragraphs[0].font.size = Pt(10)
                citation.paragraphs[0].font.color.rgb = muted
            footer_color = muted

        footer = slide.shapes.add_textbox(Inches(0.95), Inches(7.03), Inches(11.4), Inches(0.24)).text_frame
        footer.text = f"{content['title'].upper()}  ·  {index + 1:02d} / {len(content['slides']):02d}"
        footer.paragraphs[0].font.name = "Arial"; footer.paragraphs[0].font.size = Pt(10)
        footer.paragraphs[0].font.color.rgb = footer_color
        footer.paragraphs[0].alignment = PP_ALIGN.LEFT
        notes = [slide_data.get("speaker_notes", "")]
        notes.extend(f"[{source_numbers[source_id]}] {sources[source_id].title}\n{sources[source_id].url}"
                     for source_id in slide_data.get("source_ids", []) if source_id in sources)
        if slide_data.get("evidence_ids"):
            notes.append("Evidence IDs: " + ", ".join(slide_data["evidence_ids"]))
        slide.notes_slide.notes_text_frame.text = "\n\n".join(note for note in notes if note)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()
